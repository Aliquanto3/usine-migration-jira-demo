"""Vérification du Rejeu : deux Passages de même Date d'arrêté ne diffèrent que par les exclusions.

Compare fichier par fichier (`FICHIERS_PASSAGE`) : tables Parquet, JSON et JSONL champ par
champ, valeurs et liens des cellules de l'Excel, textes des slides du PowerPoint. Seuls des
chemins de `EXCLUSIONS_REJEU` peuvent différer (AD-7).

Usage : `uv run python outils/verifier_rejeu.py <passage_1> <passage_2>`
ou `uv run python outils/verifier_rejeu.py --date AAAA-MM-JJ [--racine .]`, qui compare les deux
derniers Passages publiés de la date. Code de sortie : 0 conforme, 1 non conforme, 2 usage.
"""

import argparse
import json
import sys
from datetime import date
from pathlib import Path

import polars as pl
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation

from usine.contrats.passage import EXCLUSIONS_REJEU, FICHIERS_PASSAGE

ABSENT = "<fichier>"
"""Différence signalée quand un fichier n'existe que dans un des deux Passages."""


def _differences_json(a: object, b: object, chemin: str = "") -> set[str]:
    """Chemins de champ (`etapes[].statut`) où deux valeurs JSON diffèrent."""
    if isinstance(a, dict) and isinstance(b, dict):
        return {
            d
            for cle in a.keys() | b.keys()
            for d in _differences_json(
                a.get(cle, ...), b.get(cle, ...), f"{chemin}.{cle}" if chemin else cle
            )
        }
    if isinstance(a, list) and isinstance(b, list) and len(a) == len(b):
        return {
            d for x, y in zip(a, b, strict=True) for d in _differences_json(x, y, f"{chemin}[]")
        }
    return set() if a == b else {chemin}


def _differences_cles(a: dict[str, object], b: dict[str, object]) -> set[str]:
    return {cle for cle in a.keys() | b.keys() if a.get(cle, ...) != b.get(cle, ...)}


def cellules(chemin: Path) -> dict[str, object]:
    """`<onglet>!<adresse>` → (valeur, cible d'hyperlien), plus les propriétés du classeur."""
    classeur = load_workbook(chemin)
    p = classeur.properties
    resultat: dict[str, object] = {
        "docProps": (p.created, p.modified, p.creator, p.title),
        "onglets": classeur.sheetnames,
    }
    for onglet in classeur.worksheets:
        for ligne in onglet.iter_rows():
            for c in ligne:
                resultat[f"{onglet.title}!{c.coordinate}"] = (
                    c.value,
                    c.hyperlink.target if c.hyperlink else None,
                )
    return resultat


def textes_slides(chemin: Path) -> dict[str, object]:
    """`<titre>!<libellé>` (tableaux à deux colonnes) ou `<titre>!#<forme>` → texte."""
    presentation = Presentation(str(chemin))
    p = presentation.core_properties
    textes: dict[str, object] = {"docProps": (p.created, p.modified, p.author, p.title)}

    def ajouter(cle: str, valeur: object) -> None:
        if cle in textes:
            raise ValueError(f"repère ambigu : {cle}")
        textes[cle] = valeur

    for numero, slide in enumerate(presentation.slides):
        titre_forme = slide.shapes.title
        titre = titre_forme.text_frame.text if titre_forme is not None else f"slide {numero}"
        for index, forme in enumerate(slide.shapes):
            if forme.has_table and len(forme.table.columns) == 2:
                for ligne in forme.table.rows:
                    libelle, valeur = (c.text for c in ligne.cells)
                    ajouter(f"{titre}!{libelle}", valeur)
            else:
                ajouter(
                    f"{titre}!#{index}",
                    (
                        forme.text_frame.text if forme.has_text_frame else None,
                        [[c.text for c in ligne.cells] for ligne in forme.table.rows]
                        if forme.has_table
                        else None,
                    ),
                )
    return textes


def comparer(nom: str, premier: Path, second: Path, date_arrete: str) -> set[str]:
    """Chemins qui diffèrent entre les deux Passages pour le fichier logique `nom`."""
    relatif = FICHIERS_PASSAGE[nom].format(date_arrete=date_arrete)
    a, b = premier / relatif, second / relatif
    if a.exists() != b.exists():
        return {ABSENT}
    if not a.exists():
        return set()
    if relatif.endswith(".parquet"):
        return set() if pl.read_parquet(a).equals(pl.read_parquet(b)) else {"<table>"}
    if relatif.endswith(".json"):
        return _differences_json(
            json.loads(a.read_text(encoding="utf-8")), json.loads(b.read_text(encoding="utf-8"))
        )
    if relatif.endswith(".jsonl"):
        lignes = [
            [json.loads(ligne) for ligne in p.read_text(encoding="utf-8").splitlines() if ligne]
            for p in (a, b)
        ]
        if len(lignes[0]) != len(lignes[1]):
            return {"<lignes>"}
        # Chemins depuis la racine de chaque ligne (AD-7).
        return {d for x, y in zip(*lignes, strict=True) for d in _differences_json(x, y)}
    if relatif.endswith(".xlsx"):
        return _differences_cles(cellules(a), cellules(b))
    if relatif.endswith(".pptx"):
        return _differences_cles(textes_slides(a), textes_slides(b))
    raise ValueError(f"format non comparé : {relatif}")  # pragma: no cover


def differences(premier: Path, second: Path, date_arrete: str) -> set[tuple[str, str]]:
    """(fichier logique, chemin de champ) de toutes les différences entre deux Passages."""
    return {
        (nom, d) for nom in FICHIERS_PASSAGE for d in comparer(nom, premier, second, date_arrete)
    }


def hors_exclusions(observees: set[tuple[str, str]]) -> set[tuple[str, str]]:
    """Différences que `EXCLUSIONS_REJEU` ne couvre pas : elles rendent le Rejeu non conforme."""
    return {(nom, d) for nom, d in observees if d not in EXCLUSIONS_REJEU.get(nom, ())}


def derniers_passages(racine: Path, date_arrete: date) -> tuple[Path, Path] | None:
    """Les deux derniers Passages publiés de la date (`passage_id` horodaté), ou `None`."""
    dossier = racine / "sorties" / "passages" / date_arrete.isoformat()
    if not dossier.is_dir():
        return None
    passages = sorted(p for p in dossier.iterdir() if (p / FICHIERS_PASSAGE["manifeste"]).is_file())
    if len(passages) < 2:
        return None
    return passages[-2], passages[-1]


def _date_du_passage(passage: Path) -> str:
    manifeste = json.loads((passage / FICHIERS_PASSAGE["manifeste"]).read_text(encoding="utf-8"))
    return str(manifeste["date_arrete"])


def verifier(premier: Path, second: Path) -> int:
    """Affiche les différences observées et renvoie le code de sortie (0 ou 1)."""
    date_arrete = _date_du_passage(premier)
    print(f"Rejeu du {date_arrete} : {premier.name} contre {second.name}")
    observees = differences(premier, second, date_arrete)
    fautives = hors_exclusions(observees)
    for nom, champ in sorted(observees - fautives):
        print(f"  exclue : {nom} : {champ}")
    if fautives:
        print("Rejeu NON conforme, différences hors EXCLUSIONS_REJEU :")
        for nom, champ in sorted(fautives):
            print(f"  {nom} : {champ}")
        return 1
    print(f"Rejeu conforme : {len(observees)} différence(s), toutes exclues.")
    return 0


def main(arguments: list[str] | None = None) -> int:
    """Point d'entrée en ligne de commande ; renvoie le code de sortie."""
    for flux in (sys.stdout, sys.stderr):
        # Console Windows : libellés accentués (`Traçabilité!B5`) sans erreur d'encodage.
        reconfigurer = getattr(flux, "reconfigure", None)
        if reconfigurer is not None:
            reconfigurer(encoding="utf-8", errors="replace")
    analyseur = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    analyseur.add_argument("passages", nargs="*", type=Path)
    analyseur.add_argument("--date", type=date.fromisoformat)
    analyseur.add_argument("--racine", type=Path, default=Path("."))
    args = analyseur.parse_args(arguments)
    if args.date is not None and not args.passages:
        couple = derniers_passages(args.racine, args.date)
        if couple is None:
            print(
                f"Moins de deux Passages publiés pour le {args.date.isoformat()}"
                f" sous {args.racine / 'sorties' / 'passages'}.",
                file=sys.stderr,
            )
            return 2
        return verifier(*couple)
    if args.date is None and len(args.passages) == 2:
        premier, second = args.passages
        for passage in (premier, second):
            if not (passage / FICHIERS_PASSAGE["manifeste"]).is_file():
                print(f"Passage introuvable (pas de manifeste) : {passage}", file=sys.stderr)
                return 2
        if _date_du_passage(premier) != _date_du_passage(second):
            print("Les deux Passages n'ont pas la même Date d'arrêté.", file=sys.stderr)
            return 2
        return verifier(premier, second)
    print("Donner deux Passages, ou --date AAAA-MM-JJ.", file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main())
