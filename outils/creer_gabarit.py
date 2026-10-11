"""Crée `gabarits/rapport.pptx`, gabarit PowerPoint neutre et fictif du Rapport (Story 0.6).

Gabarit 16:9 sans logo ni nom réel, construit sur le modèle par défaut de python-pptx, sans
aucune slide ; ses propriétés de document sont fixes. Usage :

    uv run python -m outils.creer_gabarit [chemin]
"""

import sys
from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

LARGEUR_16_9 = Emu(12_192_000)
"""Largeur d'une slide 16:9 (33,867 cm)."""

HAUTEUR_16_9 = Emu(6_858_000)
"""Hauteur d'une slide 16:9 (19,05 cm)."""

DATE_GABARIT = datetime(2026, 1, 1, tzinfo=UTC)
"""Date fixe des propriétés du gabarit."""

AUTEUR = "Usine de migration (démonstrateur)"

CHEMIN_PAR_DEFAUT = Path(__file__).resolve().parents[1] / "gabarits" / "rapport.pptx"


def creer_gabarit(chemin: Path) -> Path:
    """Écrit le gabarit neutre en 16:9 à `chemin` et renvoie ce chemin."""
    presentation = Presentation()
    facteur = LARGEUR_16_9 / presentation.slide_width
    presentation.slide_width = LARGEUR_16_9
    presentation.slide_height = HAUTEUR_16_9
    # Le modèle par défaut est en 4:3 : les zones du masque et des dispositions sont élargies.
    masque = presentation.slide_master
    for formes in [masque.shapes, *(disposition.shapes for disposition in masque.slide_layouts)]:
        for forme in formes:
            if forme.left is not None and forme.width is not None:
                forme.left = Emu(round(forme.left * facteur))
                forme.width = Emu(round(forme.width * facteur))
    proprietes = presentation.core_properties
    proprietes.author = AUTEUR
    proprietes.last_modified_by = AUTEUR
    proprietes.title = "Gabarit neutre du Rapport quotidien (fictif)"
    proprietes.subject = ""
    proprietes.comments = "Gabarit fictif du démonstrateur, sans logo ni nom réel."
    proprietes.keywords = ""
    proprietes.category = ""
    proprietes.revision = 1
    proprietes.created = DATE_GABARIT
    proprietes.modified = DATE_GABARIT
    proprietes.last_printed = DATE_GABARIT
    chemin.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(chemin))
    return chemin


if __name__ == "__main__":
    cible = Path(sys.argv[1]) if len(sys.argv) > 1 else CHEMIN_PAR_DEFAUT
    print(f"Gabarit écrit : {creer_gabarit(cible)}")
