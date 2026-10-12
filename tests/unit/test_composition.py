"""Racine de composition : secrets, URL Jira et sélection de l'extraction Jira (Story 2.4).

Source composite, alignements et `version_donnees` (Story 2.6).
"""

from datetime import date, timedelta
from pathlib import Path
from typing import Any

import httpx
import polars as pl
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from tests.doublures import SourceEnMemoire
from tests.unit.adaptateurs.composite.outils import (
    autre_valeur,
    etat_a_l_injection,
    injection_simulee,
    premiere_anomalie,
    transition_jira,
)
from tests.unit.adaptateurs.jira_dc.outils import ServeurFixturesDC
from tests.unit.adaptateurs.jira_dc.outils import (
    evenements_attendus as evenements_attendus_dc,
)
from tests.unit.adaptateurs.jira_http.outils import Attentes
from tests.unit.adaptateurs.journal.outils import couper_reseau, manifeste_pour
from tests.unit.chaine.environnement import (
    DATE_ARRETE,
    INSTANT,
    horloge_figee,
    horloge_pas_a_pas,
    preparer_racine,
)
from usine import composition
from usine.adaptateurs.composite import DetecteurAlignements, id_alignement
from usine.adaptateurs.entrepot_duckdb import EntrepotDuckDB, EtapeExtraction
from usine.adaptateurs.jira_cloud import SourceJiraCloud
from usine.adaptateurs.jira_dc import SourceJiraDC
from usine.adaptateurs.jira_http import construire_client_lecture
from usine.adaptateurs.journal import CHEMIN_MANIFESTE_INJECTION, SourceJournal
from usine.chaine.etapes import EtapeBouchon, EtapeMeteo, EtapePrevision, prevision_du_passage
from usine.config import charger_yaml
from usine.consultation.passage import ConsultationPassage
from usine.contrats.calendrier import TAILLE_FENETRE_GLISSANTE
from usine.contrats.config import ConfigJira, ConfigProgramme
from usine.contrats.erreurs import UsineErreur
from usine.contrats.evenements import Evenement, NatureEvenement
from usine.contrats.meteo import Couleur
from usine.contrats.passage import Etape
from usine.contrats.tickets import Champ
from usine.contrats.versions import version_donnees
from usine.domaine.meteo import DECALAGE_IND8
from usine.domaine.prevision.baseline import ParametresBaseline
from usine.domaine.prevision.risque import ParametresRisque
from usine.ports.etape import Etape as PortEtape
from usine.rendu.commun import depuis_manifeste, formater_horodatage
from usine.rendu.etape import EtapeRendu

_VARIABLES_JIRA = ("USINE_MODE_REPLI", "JIRA_EMAIL", "JIRA_TOKEN_LECTURE", "JIRA_DC_PAT")


@pytest.fixture
def racine(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    for variable in _VARIABLES_JIRA:
        monkeypatch.delenv(variable, raising=False)
    return preparer_racine(tmp_path, journal=False)


def test_url_jira(racine: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    url = charger_yaml(racine / "config" / "jira.yaml", ConfigJira).url_site
    assert composition.url_jira(racine) == url
    monkeypatch.setenv("USINE_MODE_REPLI", "1")
    assert composition.url_jira(racine) is None


def test_url_jira_sans_configuration(racine: Path) -> None:
    (racine / "config" / "jira.yaml").unlink()
    assert composition.url_jira(racine) is None


def test_secrets_illisibles(racine: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("USINE_MODE_REPLI", "peut-être")
    with pytest.raises(UsineErreur, match="USINE_MODE_REPLI"):
        composition.construire_rapport(racine)


def test_generer_journal_delegue(racine: Path) -> None:
    programme = charger_yaml(racine / "config" / "programme.yaml", ConfigProgramme)
    chemin = composition.generer_journal(programme, 42, date(2026, 10, 12), racine)
    assert chemin == racine / "donnees" / "journal.jsonl"
    assert chemin.is_file()
    assert (racine / "donnees" / "verite_terrain.json").is_file()


def test_grille_incoherente_refusee_au_chargement(racine: Path) -> None:
    """FR-18 : orange plus sévère que rouge refusé avant tout Passage (Story 3.4)."""
    chemin = racine / "config" / "grille_meteo.yaml"
    texte = chemin.read_text(encoding="utf-8")
    assert 'orange: ">= 1"' in texte
    chemin.write_text(texte.replace('orange: ">= 1"', 'orange: ">= 4"', 1), encoding="utf-8")
    with pytest.raises(UsineErreur) as erreur:
        composition.construire_rapport(racine)
    message = str(erreur.value)
    assert str(chemin) in message
    assert "IND-1" in message
    assert "« >= 4 »" in message


# --- Sélection de l'étape `extraction` (Story 2.4) ---


def _passer_en_dc(racine: Path) -> None:
    chemin = racine / "config" / "jira.yaml"
    texte = chemin.read_text(encoding="utf-8")
    assert "\nmode: cloud\n" in texte
    chemin.write_text(texte.replace("\nmode: cloud\n", "\nmode: dc\n", 1), encoding="utf-8")


def _etape_extraction(racine: Path) -> PortEtape:
    return composition.construire_rapport(racine).registre[Etape.extraction]


def _interdire_le_client(monkeypatch: pytest.MonkeyPatch) -> None:
    def construire(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("client Jira construit sans extraction attendue")

    monkeypatch.setattr(composition, "construire_client_lecture", construire)


@pytest.fixture
def racine_journal(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    for variable in _VARIABLES_JIRA:
        monkeypatch.delenv(variable, raising=False)
    return preparer_racine(tmp_path)


def test_sans_secret_extraction_bouchon(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _interdire_le_client(monkeypatch)
    assert isinstance(_etape_extraction(racine_journal), EtapeBouchon)
    # Le secret de l'autre plateforme ne suffit pas.
    monkeypatch.setenv("JIRA_DC_PAT", "pat-fictif")
    assert isinstance(_etape_extraction(racine_journal), EtapeBouchon)


@pytest.mark.parametrize("mode", ["cloud", "dc"])
def test_secret_blanc_extraction_bouchon(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch, mode: str
) -> None:
    if mode == "dc":
        _passer_en_dc(racine_journal)
    monkeypatch.setenv("JIRA_DC_PAT", "   ")
    monkeypatch.setenv("JIRA_TOKEN_LECTURE", "   ")
    monkeypatch.setenv("JIRA_EMAIL", "lecteur@example.invalid")
    _interdire_le_client(monkeypatch)
    assert isinstance(_etape_extraction(racine_journal), EtapeBouchon)


def test_sans_configuration_jira_extraction_bouchon(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (racine_journal / "config" / "jira.yaml").unlink()
    monkeypatch.setenv("JIRA_DC_PAT", "pat-fictif")
    monkeypatch.setenv("JIRA_TOKEN_LECTURE", "jeton-fictif")
    _interdire_le_client(monkeypatch)
    assert isinstance(_etape_extraction(racine_journal), EtapeBouchon)


def test_mode_repli_coupe_l_extraction(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _passer_en_dc(racine_journal)
    monkeypatch.setenv("JIRA_DC_PAT", "pat-fictif")
    monkeypatch.setenv("USINE_MODE_REPLI", "1")
    _interdire_le_client(monkeypatch)
    assert isinstance(_etape_extraction(racine_journal), EtapeBouchon)


def test_cloud_avec_jeton_et_email(racine_journal: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("JIRA_TOKEN_LECTURE", "jeton-fictif")
    monkeypatch.setenv("JIRA_EMAIL", "lecteur@example.invalid")
    etape = _etape_extraction(racine_journal)
    assert isinstance(etape, EtapeExtraction)
    assert isinstance(etape.source, SourceJiraCloud)
    assert etape.chemin == racine_journal / "donnees" / "entrepot.duckdb"
    assert etape.marge == timedelta(minutes=10)


def test_dc_avec_pat(racine_journal: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _passer_en_dc(racine_journal)
    monkeypatch.setenv("JIRA_DC_PAT", "pat-fictif")
    etape = _etape_extraction(racine_journal)
    assert isinstance(etape, EtapeExtraction)
    assert isinstance(etape.source, SourceJiraDC)
    assert etape.chemin == racine_journal / "donnees" / "entrepot.duckdb"
    marge = charger_yaml(
        racine_journal / "config" / "jira.yaml", ConfigJira
    ).marge_extraction_minutes
    assert etape.marge == timedelta(minutes=marge)


def test_cloud_sans_email_nomme_la_variable(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("JIRA_TOKEN_LECTURE", "jeton-fictif")
    with pytest.raises(UsineErreur, match="JIRA_EMAIL"):
        composition.construire_rapport(racine_journal)


def test_passage_sans_secret_journal_seul_sans_requete(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _interdire_le_client(monkeypatch)
    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    assert lance.manifeste.source == "journal"
    assert lance.repertoire.is_dir()
    assert not (racine_journal / "donnees" / "entrepot.duckdb").exists()


def test_registre_prevision_et_rendu(racine_journal: Path) -> None:
    """Story 5.3 : `prevision` et `rendu` au registre ; Prévisions des Entités rouges publiées,
    tracées à l'horloge ; contexte du rendu relu du manifeste."""
    rapport = composition.construire_rapport(racine_journal, horloge=horloge_figee())
    assert isinstance(rapport.registre[Etape.prevision], EtapePrevision)
    rendu = rapport.registre[Etape.rendu]
    assert isinstance(rendu, EtapeRendu)
    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    consultation = ConsultationPassage(lance.repertoire)
    previsions = consultation.previsions()
    assert [p.entite for p in previsions] == [
        m.entite for m in consultation.meteo() if m.couleur is Couleur.rouge
    ]
    assert {t.horodatage for t in consultation.traces()} == {INSTANT}
    contexte = depuis_manifeste(lance.manifeste, rapport.configuration.grille_meteo)
    assert (contexte.source, contexte.version_grille, contexte.version_donnees) == (
        rapport.donnees.source,
        rapport.donnees.version_grille,
        rapport.donnees.version_donnees,
    )
    assert contexte.previsions_calculees


def _capturer_prevision(monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    capture: dict[str, Any] = {}
    origine = prevision_du_passage

    def prevision(*args: Any, **kwargs: Any) -> Any:
        capture.update(kwargs)
        return origine(*args, **kwargs)

    monkeypatch.setattr(composition, "prevision_du_passage", prevision)
    return capture


def test_parametres_de_chaine_yaml(racine_journal: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Story 5.4 : la Prévision reçoit `prevision` de `chaine.yaml`, `EtapeMeteo` le décalage."""
    (racine_journal / "config" / "chaine.yaml").write_text(
        "palier: 1\n"
        "prevision:\n"
        "  baseline:\n"
        "    horizon: 7\n"
        "    fenetre_velocite: 8\n"
        "  risque:\n"
        "    fenetre: 6\n"
        "    horizon: 3\n"
        "decalage_comparaison: 12\n",
        encoding="utf-8",
    )
    capture = _capturer_prevision(monkeypatch)
    rapport = composition.construire_rapport(racine_journal, horloge=horloge_figee())
    assert capture["baseline"] == ParametresBaseline(horizon=7, fenetre_velocite=8)
    assert capture["risque"] == ParametresRisque(fenetre=6, horizon=3)
    meteo = rapport.registre[Etape.meteo]
    assert isinstance(meteo, EtapeMeteo)
    assert meteo.taille_historique == max(12, DECALAGE_IND8, TAILLE_FENETRE_GLISSANTE) + 1


def test_parametres_par_defaut(racine_journal: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Un `chaine.yaml` réduit à `palier` donne les défauts du domaine (Passage inchangé)."""
    (racine_journal / "config" / "chaine.yaml").write_text("palier: 1\n", encoding="utf-8")
    capture = _capturer_prevision(monkeypatch)
    rapport = composition.construire_rapport(racine_journal, horloge=horloge_figee())
    assert capture["baseline"] == ParametresBaseline()
    assert capture["risque"] == ParametresRisque()
    meteo = rapport.registre[Etape.meteo]
    assert isinstance(meteo, EtapeMeteo)
    assert meteo.taille_historique == max(DECALAGE_IND8, TAILLE_FENETRE_GLISSANTE) + 1


# --- Journal lu aux clés Jira (Story 2.5) ---


def _cles_tickets_etats(repertoire: Path) -> set[str]:
    return set(pl.read_parquet(repertoire / "tickets_etats.parquet")["cle_ticket"].to_list())


def test_passage_journal_seul_avec_manifeste_aux_cles_jira(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    brut = SourceJournal(racine_journal / "donnees" / "journal.jsonl")
    manifeste = manifeste_pour(brut.entete, brut.evenements())
    chemin = racine_journal / CHEMIN_MANIFESTE_INJECTION
    chemin.parent.mkdir(parents=True)
    chemin.write_text(manifeste.model_dump_json(indent=2), encoding="utf-8")
    _interdire_le_client(monkeypatch)
    couper_reseau(monkeypatch)
    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    assert lance.manifeste.source == "journal"
    cles = _cles_tickets_etats(lance.repertoire)
    assert cles
    assert cles <= set(manifeste.correspondance_cles.values())
    assert not cles & set(manifeste.correspondance_cles)


def test_passage_sans_manifeste_aux_cles_generees(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    brut = SourceJournal(racine_journal / "donnees" / "journal.jsonl")
    _interdire_le_client(monkeypatch)
    couper_reseau(monkeypatch)
    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    cles = _cles_tickets_etats(lance.repertoire)
    assert cles
    assert cles <= {e.cle_ticket for e in brut.evenements()}


def test_manifeste_invalide_nomme_le_fichier(racine_journal: Path) -> None:
    chemin = racine_journal / CHEMIN_MANIFESTE_INJECTION
    chemin.parent.mkdir(parents=True)
    chemin.write_text("[]", encoding="utf-8")
    with pytest.raises(UsineErreur, match=r"manifeste_injection.json"):
        composition.construire_rapport(racine_journal)


def test_passage_dc_de_bout_en_bout(racine_journal: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _passer_en_dc(racine_journal)
    monkeypatch.setenv("JIRA_DC_PAT", "pat-fictif")
    serveur = ServeurFixturesDC()
    fermetures: list[str] = []

    def construire_simule(*args: Any, **kwargs: Any) -> Any:
        client = construire_client_lecture(
            *args, **kwargs, attendre=Attentes(), transport_http=httpx.MockTransport(serveur)
        )
        fermer = client.fermer

        def fermer_espion() -> None:
            fermetures.append("fermé")
            fermer()

        monkeypatch.setattr(client, "fermer", fermer_espion)
        return client

    monkeypatch.setattr(composition, "construire_client_lecture", construire_simule)
    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    assert lance.repertoire.is_dir()
    assert lance.manifeste.source == "journal"
    assert serveur.inattendues == []
    assert {r.headers["Authorization"] for r in serveur.requetes} == {"Bearer pat-fictif"}
    assert fermetures == ["fermé"]
    entrepot = EntrepotDuckDB(racine_journal / "donnees" / "entrepot.duckdb")
    assert tuple(entrepot.evenements()) == evenements_attendus_dc()
    assert entrepot.repere() == INSTANT


def test_client_ferme_si_le_passage_echoue_avant_l_extraction(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _passer_en_dc(racine_journal)
    monkeypatch.setenv("JIRA_DC_PAT", "pat-fictif")
    serveur = ServeurFixturesDC()
    fermetures: list[str] = []

    def construire_simule(*args: Any, **kwargs: Any) -> Any:
        client = construire_client_lecture(
            *args, **kwargs, attendre=Attentes(), transport_http=httpx.MockTransport(serveur)
        )
        monkeypatch.setattr(client, "fermer", lambda: fermetures.append("fermé"))
        return client

    def echouer(*args: Any, **kwargs: Any) -> Any:
        raise UsineErreur("Passage avorté avant l'extraction")

    monkeypatch.setattr(composition, "construire_client_lecture", construire_simule)
    monkeypatch.setattr(composition, "executer_passage", echouer)
    with pytest.raises(UsineErreur, match="avorté"):
        composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    assert fermetures == ["fermé"]
    assert serveur.requetes == []


# --- Source composite (Story 2.6) ---


def _ecrire_manifeste(racine: Path) -> SourceJournal:
    """Écrit un manifeste d'injection fictif et rend le Journal lu aux clés Jira."""
    chemin_journal = racine / "donnees" / "journal.jsonl"
    brut = SourceJournal(chemin_journal)
    manifeste = manifeste_pour(brut.entete, brut.evenements())
    chemin = racine / CHEMIN_MANIFESTE_INJECTION
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_text(manifeste.model_dump_json(indent=2), encoding="utf-8")
    return SourceJournal(chemin_journal, manifeste)


class _ClientDouble:
    """Client Jira doublé : aucune requête, fermeture notée."""

    def __init__(self, fermetures: list[str]) -> None:
        self._fermetures = fermetures

    def fermer(self) -> None:
        self._fermetures.append("fermé")


def _doubler_jira(
    monkeypatch: pytest.MonkeyPatch, evenements: list[Evenement], fermetures: list[str]
) -> None:
    """Cloud actif (jeton et e-mail fictifs) ; client et source Jira remplacés par des doublures."""
    monkeypatch.setenv("JIRA_TOKEN_LECTURE", "jeton-fictif")
    monkeypatch.setenv("JIRA_EMAIL", "lecteur@example.invalid")
    monkeypatch.setattr(
        composition, "construire_client_lecture", lambda *a, **k: _ClientDouble(fermetures)
    )
    monkeypatch.setattr(composition, "SourceJiraCloud", lambda *a, **k: SourceEnMemoire(evenements))
    couper_reseau(monkeypatch)


def _etat(repertoire: Path, cle: str, colonne: str, jour: date = DATE_ARRETE) -> object:
    etats = pl.read_parquet(repertoire / "tickets_etats.parquet")
    ligne = etats.filter((pl.col("date") == jour) & (pl.col("cle_ticket") == cle))
    return ligne[colonne].item()


def _alignements(racine: Path) -> list[tuple[str, object]]:
    entrepot = EntrepotDuckDB(racine / "donnees" / "entrepot.duckdb")
    return [
        (e.id_evenement, e.horodatage)
        for e in entrepot.evenements()
        if e.nature is NatureEvenement.alignement
    ]


def test_passage_composite_de_bout_en_bout(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    journal = _ecrire_manifeste(racine_journal)
    assert journal.manifeste is not None
    fin = journal.manifeste.horodatage_fin_injection
    ecart = premiere_anomalie(journal)
    gravite = autre_valeur(Champ.gravite, etat_a_l_injection(journal, ecart, Champ.gravite))
    manuel = premiere_anomalie(journal, sauf=frozenset({ecart}))
    statut = etat_a_l_injection(journal, manuel, Champ.statut)
    nouveau_statut = autre_valeur(Champ.statut, statut)
    jira = [
        *injection_simulee(journal, {ecart: {Champ.gravite: gravite}}),
        transition_jira(manuel, statut, nouveau_statut, fin + timedelta(hours=1)),
    ]
    fermetures: list[str] = []
    _doubler_jira(monkeypatch, jira, fermetures)

    rapport = composition.construire_rapport(racine_journal, horloge=horloge_figee())
    extraction = rapport.registre[Etape.extraction]
    assert isinstance(extraction, EtapeExtraction)
    assert isinstance(extraction.aligner, DetecteurAlignements)
    assert rapport.donnees.source == "composite"
    assert rapport.horodatage_extraction == INSTANT
    extraction.fermer()

    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())

    assert lance.manifeste.source == "composite"
    assert lance.manifeste.horodatage_extraction == INSTANT
    assert lance.manifeste.nb_alignements == 1
    assert fermetures == ["fermé", "fermé"]
    # La transition faite à la main dans Jira après l'injection apparaît au Passage.
    assert statut != nouveau_statut
    assert _etat(lance.repertoire, manuel, "statut") == nouveau_statut
    # Le classeur publié affiche la source et l'instant d'extraction transmis au rendu.
    classeur = load_workbook(lance.repertoire / "rapport-2026-10-09.xlsx")
    tracabilite = {
        ligne[0]: ligne[1] for ligne in classeur["Traçabilité"].iter_rows(values_only=True)
    }
    assert tracabilite["Source"] == "composite"
    assert lance.manifeste.horodatage_extraction == INSTANT
    assert tracabilite["Horodatage d'extraction"] == formater_horodatage(INSTANT)
    # L'alignement, horodaté à la détection (après la Date d'arrêté), ne change pas cette date.
    assert _etat(lance.repertoire, ecart, "gravite") != gravite
    attendu = [(id_alignement(ecart, Champ.gravite, gravite), INSTANT)]
    assert _alignements(racine_journal) == attendu
    entrepot = EntrepotDuckDB(racine_journal / "donnees" / "entrepot.duckdb")
    assert entrepot.repere() == INSTANT

    # Deuxième Passage, même écart : rien d'ajouté, première détection gardée.
    plus_tard = INSTANT + timedelta(days=1)
    second = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee(plus_tard))
    assert second.manifeste.nb_alignements == 0
    assert second.manifeste.horodatage_extraction == plus_tard
    assert _alignements(racine_journal) == attendu
    assert entrepot.repere() == plus_tard
    assert _etat(second.repertoire, manuel, "statut") == nouveau_statut


def test_passage_composite_instant_d_extraction_partage(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Horloge qui avance : repère, alignement, manifeste et empreinte gardent un même instant.

    Passage à la Date d'arrêté du jour de la détection : la valeur Jira alignée y est publiée.
    """
    journal = _ecrire_manifeste(racine_journal)
    ecart = premiere_anomalie(journal)
    gravite = autre_valeur(Champ.gravite, etat_a_l_injection(journal, ecart, Champ.gravite))
    _doubler_jira(monkeypatch, injection_simulee(journal, {ecart: {Champ.gravite: gravite}}), [])
    jour_detection = date(2026, 10, 12)
    assert jour_detection == INSTANT.date()

    lance = composition.lancer_rapport(racine_journal, jour_detection, horloge_pas_a_pas())

    instant = lance.manifeste.horodatage_extraction
    assert instant is not None
    assert instant != lance.manifeste.horodatage_debut  # l'horloge a bien avancé
    entrepot = EntrepotDuckDB(racine_journal / "donnees" / "entrepot.duckdb")
    assert entrepot.repere() == instant
    assert _alignements(racine_journal) == [(id_alignement(ecart, Champ.gravite, gravite), instant)]
    octets = (racine_journal / "donnees" / "journal.jsonl").read_bytes()
    octets += b"\n" + (racine_journal / CHEMIN_MANIFESTE_INJECTION).read_bytes()
    assert lance.manifeste.version_donnees == version_donnees(octets, instant)
    # Jira fait foi : à la Date d'arrêté de la détection, la gravité publiée est celle de Jira.
    assert _etat(lance.repertoire, ecart, "gravite", jour_detection) == gravite
    assert _etat(lance.repertoire, ecart, "gravite") != gravite


def test_extraction_sans_manifeste_reste_journal(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _doubler_jira(monkeypatch, [], [])
    rapport = composition.construire_rapport(racine_journal, horloge=horloge_figee())
    extraction = rapport.registre[Etape.extraction]
    assert isinstance(extraction, EtapeExtraction)
    assert extraction.aligner is None
    assert rapport.donnees.source == "journal"
    assert rapport.horodatage_extraction is None
    lance = composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee())
    assert lance.manifeste.source == "journal"
    assert lance.manifeste.horodatage_extraction is None
    assert lance.manifeste.nb_alignements == 0


def test_detection_avant_la_fin_d_injection_echoue(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    journal = _ecrire_manifeste(racine_journal)
    assert journal.manifeste is not None
    fin = journal.manifeste.horodatage_fin_injection
    _doubler_jira(monkeypatch, injection_simulee(journal), [])
    with pytest.raises(UsineErreur, match="fin d'injection") as erreur:
        composition.lancer_rapport(racine_journal, DATE_ARRETE, horloge_figee(fin))
    assert "extraction" in str(erreur.value)
    entrepot = EntrepotDuckDB(racine_journal / "donnees" / "entrepot.duckdb")
    assert entrepot.repere() is None


def test_version_donnees_trois_empreintes(
    racine_journal: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    octets_journal = (racine_journal / "donnees" / "journal.jsonl").read_bytes()
    sans_manifeste = composition.construire_rapport(racine_journal).donnees.version_donnees
    assert sans_manifeste == version_donnees(octets_journal, None)

    _ecrire_manifeste(racine_journal)
    octets = octets_journal + b"\n" + (racine_journal / CHEMIN_MANIFESTE_INJECTION).read_bytes()
    avec_manifeste = composition.construire_rapport(racine_journal).donnees.version_donnees
    assert avec_manifeste == version_donnees(octets, None)

    _doubler_jira(monkeypatch, [], [])
    rapport = composition.construire_rapport(racine_journal, horloge=horloge_figee())
    assert rapport.donnees.version_donnees == version_donnees(octets, INSTANT)
    assert len({sans_manifeste, avec_manifeste, rapport.donnees.version_donnees}) == 3
