"""Recette du Rapport sur le Journal complet (Story 5.6, AC 1).

Test `lent` : exclu par défaut (`pyproject.toml`), lancé par `uv run pytest -m lent`. Un Passage
Palier 1 (aucun appel LLM) s'exécute par `lancer_rapport` sur le Journal généré (graine 42,
génération le 2026-10-12) avec `config/programme.yaml` livré (10 Entités, volumétrie cible).
À la Date d'arrêté 2026-10-09 : A et B rouges, C orange, une slide pour chacune, Prévision de
sortie du rouge pour A, aucune sortie projetée pour B, Risque d'entrée en rouge sur C.
"""

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from tests.unit.chaine.environnement import DATE_ARRETE, horloge_figee, preparer_racine
from tests.unit.rendu.test_powerpoint import _lire
from usine.composition import lancer_rapport
from usine.consultation.passage import ConsultationPassage
from usine.contrats.meteo import Couleur
from usine.contrats.prevision import StatutPrevision
from usine.rendu.commun import texte_prevision

pytestmark = pytest.mark.lent

ONGLETS = [
    "Synthèse",
    "Météo détaillée",
    "KPI",
    "Tendances",
    "Anomalies ouvertes",
    "Tests",
    "Certifications",
    "Commentaires",
    "Traçabilité",
]


def test_rapport_sur_le_journal_complet(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    racine = preparer_racine(tmp_path, programme_complet=True)
    monkeypatch.chdir(racine)  # aucun `.env` du poste lu par `Secrets`
    monkeypatch.delenv("USINE_MODE_REPLI", raising=False)
    repertoire = lancer_rapport(racine, DATE_ARRETE, horloge_figee(), lambda: "5e6000").repertoire

    consultation = ConsultationPassage(repertoire)
    manifeste = consultation.manifeste()
    assert manifeste.palier == 1
    assert manifeste.date_arrete == DATE_ARRETE

    couleurs = {meteo.entite: meteo.couleur for meteo in consultation.meteo()}
    assert list(couleurs) == [f"ENT-{lettre}" for lettre in "ABCDEFGHIJ"]
    assert couleurs["ENT-A"] is Couleur.rouge
    assert couleurs["ENT-B"] is Couleur.rouge
    assert couleurs["ENT-C"] is Couleur.orange

    statuts = {p.entite: p.statut for p in consultation.previsions()}
    assert statuts["ENT-A"] is StatutPrevision.sortie_projetee
    assert statuts["ENT-B"] is StatutPrevision.pas_de_sortie_projetee
    assert consultation.risques(entite="ENT-C")

    nom = f"rapport-{DATE_ARRETE.isoformat()}"
    classeur = load_workbook(io.BytesIO((repertoire / f"{nom}.xlsx").read_bytes()))
    assert classeur.sheetnames == ONGLETS

    slides = _lire((repertoire / f"{nom}.pptx").read_bytes())
    titres = [slide.titre for slide in slides]
    rouges_oranges = [
        f"{m.entite} — {'Rouge' if m.couleur is Couleur.rouge else 'Orange'}"
        for m in consultation.meteo()
        if m.couleur in (Couleur.rouge, Couleur.orange)
    ]
    assert rouges_oranges[:3] == ["ENT-A — Rouge", "ENT-B — Rouge", "ENT-C — Orange"]
    assert titres[3 : 3 + len(rouges_oranges)] == rouges_oranges
    assert titres[3:6] == ["ENT-A — Rouge", "ENT-B — Rouge", "ENT-C — Orange"]
    for slide, entite in zip(slides[3:6], ("ENT-A", "ENT-B", "ENT-C"), strict=True):
        texte = texte_prevision(consultation, entite, DATE_ARRETE, True)
        assert slide.zones[2] == f"Prévision\n{texte}"
        assert texte.strip()
