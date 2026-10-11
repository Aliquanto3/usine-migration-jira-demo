"""Commande `usine rapport [--date-arrete]` (Story 0.5)."""

import json
from dataclasses import replace
from datetime import date, datetime
from pathlib import Path

import polars as pl
import pytest
import time_machine
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation

from tests.unit.chaine.environnement import preparer_racine
from usine import cli, composition
from usine.chaine.cli import lire_date_arrete
from usine.contrats.calendrier import FUSEAU
from usine.contrats.erreurs import UsineErreur
from usine.contrats.passage import EXCLUSIONS_REJEU, ContextePassage, Etape, ResultatEtape
from usine.ports.consultation import Consultation

VERROU = Path("donnees") / ".verrou-passage"


def _usine(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str], *arguments: str
) -> tuple[int, str, str]:
    monkeypatch.setattr("sys.argv", ["usine", "rapport", *arguments])
    with pytest.raises(SystemExit) as sortie:
        cli.main()
    code = sortie.value.code
    sorties = capsys.readouterr()
    return (0 if code is None else int(code)), sorties.out, sorties.err


@pytest.fixture
def racine(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    preparer_racine(tmp_path)
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _passages(racine: Path) -> list[Path]:
    return sorted((racine / "sorties" / "passages" / "2026-10-09").iterdir())


def _sans_exclusions(valeur: object) -> object:
    if isinstance(valeur, dict):
        return {
            cle: _sans_exclusions(v) for cle, v in valeur.items() if cle not in EXCLUSIONS_REJEU
        }
    if isinstance(valeur, list):
        return [_sans_exclusions(v) for v in valeur]
    return valeur


def _excel_relu(chemin: Path) -> object:
    """Valeurs, hyperliens et propriétés relus de l'Excel (les octets du zip varient)."""
    classeur = load_workbook(chemin)
    p = classeur.properties
    return (
        classeur.sheetnames,
        [
            [(c.value, c.hyperlink.target if c.hyperlink else None) for c in ligne]
            for onglet in classeur.worksheets
            for ligne in onglet.iter_rows()
        ],
        (p.created, p.modified, p.creator, p.title),
    )


def _powerpoint_relu(chemin: Path) -> object:
    """Textes, tableaux et propriétés relus du PowerPoint."""
    presentation = Presentation(str(chemin))
    p = presentation.core_properties
    formes = [
        (
            forme.text_frame.text if forme.has_text_frame else None,
            [[c.text for c in ligne.cells] for ligne in forme.table.rows]
            if forme.has_table
            else None,
        )
        for slide in presentation.slides
        for forme in slide.shapes
    ]
    return formes, (p.created, p.modified, p.author, p.title)


def test_passage_publie_et_rejeu(
    racine: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    code, sortie, _ = _usine(monkeypatch, capsys, "--date-arrete", "2026-10-09")
    assert code == 0
    assert "Passage publié" in sortie
    code, _, _ = _usine(monkeypatch, capsys, "--date-arrete", "2026-10-09")
    assert code == 0
    premier, second = _passages(racine)
    assert premier.name != second.name
    assert not (racine / VERROU).exists()
    assert not list((racine / "sorties" / "passages").glob(".en-cours-*"))
    for nom in ("meteo.json", "manifeste.json"):
        a, b = (json.loads((p / nom).read_text(encoding="utf-8")) for p in (premier, second))
        assert _sans_exclusions(a) == _sans_exclusions(b), nom
    for nom in ("tickets_etats.parquet", "instantanes.parquet", "indicateurs.parquet"):
        assert pl.read_parquet(premier / nom).equals(pl.read_parquet(second / nom)), nom
    xlsx, pptx = "rapport-2026-10-09.xlsx", "rapport-2026-10-09.pptx"
    assert _excel_relu(premier / xlsx) == _excel_relu(second / xlsx)
    assert _powerpoint_relu(premier / pptx) == _powerpoint_relu(second / pptx)


@pytest.mark.parametrize(
    ("valeur", "attendu"),
    [("2026-10-10", "2026-10-10"), ("2026-13-01", "2026-13-01"), ("9/10/2026", "9/10/2026")],
)
def test_date_refusee(
    racine: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    valeur: str,
    attendu: str,
) -> None:
    code, _, erreur = _usine(monkeypatch, capsys, "--date-arrete", valeur)
    assert code == 1
    assert attendu in erreur
    assert not (racine / "sorties").exists()
    assert not (racine / VERROU).exists()


def test_date_par_defaut() -> None:
    lundi = datetime(2026, 10, 12, 0, 30, tzinfo=FUSEAU)  # encore dimanche en UTC
    assert lire_date_arrete(None, lundi).isoformat() == "2026-10-09"
    with time_machine.travel(lundi, tick=False):
        assert lire_date_arrete(None, datetime.now(FUSEAU)).isoformat() == "2026-10-09"
    assert lire_date_arrete("2026-10-06", lundi).isoformat() == "2026-10-06"


def test_date_par_defaut_en_cli(
    racine: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    with time_machine.travel(datetime(2026, 10, 12, 10, tzinfo=FUSEAU), tick=False):
        code, _, erreur = _usine(monkeypatch, capsys)
    assert code == 0, erreur
    assert len(_passages(racine)) == 1


def test_verrou_pris(
    racine: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    (racine / VERROU).write_text("autre\n", encoding="utf-8")
    code, _, erreur = _usine(monkeypatch, capsys, "--date-arrete", "2026-10-09")
    assert code == 1
    assert ".verrou-passage" in erreur
    assert (racine / VERROU).read_text(encoding="utf-8") == "autre\n"
    assert not (racine / "sorties").exists()


def test_journal_absent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    preparer_racine(tmp_path, journal=False)
    monkeypatch.chdir(tmp_path)
    code, _, erreur = _usine(monkeypatch, capsys, "--date-arrete", "2026-10-09")
    assert code == 1
    assert "journal.jsonl" in erreur
    assert not (tmp_path / VERROU).exists()


class _Panne:
    def executer(self, ctx: ContextePassage[Consultation]) -> ResultatEtape:
        raise RuntimeError("panne")


def test_echec_rend_le_verrou(
    racine: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    construire = composition.construire_rapport

    def construire_en_panne(r: Path) -> composition.Rapport:
        rapport = construire(r)
        return replace(rapport, registre={**rapport.registre, Etape.kpi: _Panne()})

    monkeypatch.setattr(composition, "construire_rapport", construire_en_panne)
    code, _, erreur = _usine(monkeypatch, capsys, "--date-arrete", "2026-10-09")
    assert code == 1
    assert "kpi" in erreur
    assert "panne" in erreur
    assert not (racine / VERROU).exists()
    assert len(list((racine / "sorties" / "echecs").iterdir())) == 1
    assert not (racine / "sorties" / "passages" / "2026-10-09").exists()


def test_lire_date_arrete_refus() -> None:
    maintenant = datetime(2026, 10, 12, 10, tzinfo=FUSEAU)
    with pytest.raises(UsineErreur, match="2026-10-11"):
        lire_date_arrete("2026-10-11", maintenant)
    with pytest.raises(UsineErreur, match="2026-02-30"):
        lire_date_arrete("2026-02-30", maintenant)


def test_lancer_rapport_refuse_un_jour_non_ouvre(racine: Path) -> None:
    with pytest.raises(UsineErreur, match="2026-10-10") as erreur:
        composition.lancer_rapport(racine, date(2026, 10, 10))
    assert "ce n'est pas un Jour ouvré" in str(erreur.value)
    assert not (racine / "sorties").exists()
    assert not (racine / VERROU).exists()
