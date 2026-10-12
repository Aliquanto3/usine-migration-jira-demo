"""Rejeu idempotent sur source composite simulée, échec sans rapport, journal de progression (5.5).

Deux Passages Palier 1 de même Date d'arrêté, dont l'instant d'extraction et la version des
données diffèrent, sont comparés fichier par fichier (`FICHIERS_PASSAGE`) : seuls des chemins
de `EXCLUSIONS_REJEU` peuvent différer, et les champs de traçabilité exclus diffèrent bien.
"""

import json
import logging
from collections.abc import Iterator
from dataclasses import replace
from datetime import timedelta
from pathlib import Path
from typing import Any

import polars as pl
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation

from tests.unit.chaine.environnement import INSTANT, preparer_racine
from usine import cli, composition
from usine.contrats.passage import (
    EXCLUSIONS_REJEU,
    FICHIERS_PASSAGE,
    ContextePassage,
    DonneesPassage,
    Etape,
    ResultatEtape,
    StatutEtape,
)
from usine.contrats.versions import version_donnees
from usine.ports.consultation import Consultation
from usine.rendu.etape import EtapeRendu

DATE = "2026-10-09"
DIFFERENCES_ATTENDUES = {
    ("manifeste", "horodatage_extraction"),
    ("manifeste", "version_donnees"),
    ("manifeste", "passage_id"),
    ("meteo", "passage_id"),
    ("previsions", "passage_id"),
    ("trace_previsions", "passage_id"),
    ("rapport_xlsx", "Traçabilité!B5"),
    ("rapport_xlsx", "Traçabilité!B8"),
    ("rapport_pptx", "Annexe méthode!Horodatage d'extraction"),
    ("rapport_pptx", "Annexe méthode!Version des données"),
}
"""Différences qui doivent effectivement apparaître : sinon l'exclusion ne serait pas exercée."""


@pytest.fixture
def racine(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    preparer_racine(tmp_path)
    monkeypatch.chdir(tmp_path)
    return tmp_path


def _usine(monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> tuple[int, str]:
    monkeypatch.setattr("sys.argv", ["usine", "rapport", "--date-arrete", DATE])
    with pytest.raises(SystemExit) as sortie:
        cli.main()
    code = sortie.value.code
    return (0 if code is None else int(code)), capsys.readouterr().err


def _simuler_composite(monkeypatch: pytest.MonkeyPatch) -> None:
    """Chaque composition reçoit son propre instant d'extraction et sa version des données."""
    construire = composition.construire_rapport
    instants: Iterator[int] = iter(range(1, 100))

    def construire_composite(r: Path, *args: Any, **kwargs: Any) -> composition.Rapport:
        rapport = construire(r, *args, **kwargs)
        horodatage = INSTANT + timedelta(hours=next(instants))
        donnees = DonneesPassage(
            "composite",
            rapport.donnees.version_grille,
            version_donnees(b"journal composite", horodatage),
        )
        gabarit = (r / composition.CHEMIN_GABARIT).read_bytes()
        registre = {**rapport.registre, Etape.rendu: EtapeRendu(gabarit, donnees, horodatage)}
        return replace(
            rapport, donnees=donnees, horodatage_extraction=horodatage, registre=registre
        )

    monkeypatch.setattr(composition, "construire_rapport", construire_composite)


def _differences(a: object, b: object, chemin: str = "") -> set[str]:
    """Chemins de champ (`etapes[].statut`) où deux valeurs JSON diffèrent."""
    if isinstance(a, dict) and isinstance(b, dict):
        return {
            d
            for cle in a.keys() | b.keys()
            for d in _differences(
                a.get(cle, ...), b.get(cle, ...), f"{chemin}.{cle}" if chemin else cle
            )
        }
    if isinstance(a, list) and isinstance(b, list) and len(a) == len(b):
        return {d for x, y in zip(a, b, strict=True) for d in _differences(x, y, f"{chemin}[]")}
    return set() if a == b else {chemin}


def _differences_cles(a: dict[str, object], b: dict[str, object]) -> set[str]:
    return {cle for cle in a.keys() | b.keys() if a.get(cle, ...) != b.get(cle, ...)}


def _cellules(chemin: Path) -> dict[str, object]:
    """`<onglet>!<adresse>` → (valeur, cible d'hyperlien), plus les propriétés du classeur."""
    classeur = load_workbook(chemin)
    p = classeur.properties
    cellules: dict[str, object] = {
        "docProps": (p.created, p.modified, p.creator, p.title),
        "onglets": classeur.sheetnames,
    }
    for onglet in classeur.worksheets:
        for ligne in onglet.iter_rows():
            for c in ligne:
                cellules[f"{onglet.title}!{c.coordinate}"] = (
                    c.value,
                    c.hyperlink.target if c.hyperlink else None,
                )
    return cellules


def _textes_slides(chemin: Path) -> dict[str, object]:
    """`<titre>!<libellé>` (tableaux à deux colonnes) ou `<titre>!#<forme>` → texte."""
    presentation = Presentation(str(chemin))
    p = presentation.core_properties
    textes: dict[str, object] = {"docProps": (p.created, p.modified, p.author, p.title)}

    def ajouter(cle: str, valeur: object) -> None:
        assert cle not in textes, f"repère ambigu : {cle}"
        textes[cle] = valeur

    for numero, slide in enumerate(presentation.slides):
        titre_forme = slide.shapes.title
        titre = titre_forme.text_frame.text if titre_forme is not None else f"slide {numero}"
        for index, forme in enumerate(slide.shapes):
            if forme.has_table and len(forme.table.columns) == 2:
                for ligne in forme.table.rows:
                    libelle, valeur = (c.text for c in ligne.cells)
                    ajouter(f"{titre}!{libelle}", valeur)
            else:
                ajouter(
                    f"{titre}!#{index}",
                    (
                        forme.text_frame.text if forme.has_text_frame else None,
                        [[c.text for c in ligne.cells] for ligne in forme.table.rows]
                        if forme.has_table
                        else None,
                    ),
                )
    return textes


def _comparer(nom: str, premier: Path, second: Path) -> set[str]:
    """Chemins qui diffèrent entre les deux Passages pour le fichier logique `nom`."""
    relatif = FICHIERS_PASSAGE[nom].format(date_arrete=DATE)
    a, b = premier / relatif, second / relatif
    assert a.exists() == b.exists(), nom
    if not a.exists():
        return set()
    if relatif.endswith(".parquet"):
        return set() if pl.read_parquet(a).equals(pl.read_parquet(b)) else {"<table>"}
    if relatif.endswith(".json"):
        return _differences(
            json.loads(a.read_text(encoding="utf-8")), json.loads(b.read_text(encoding="utf-8"))
        )
    if relatif.endswith(".jsonl"):
        lignes = [
            [json.loads(ligne) for ligne in p.read_text(encoding="utf-8").splitlines() if ligne]
            for p in (a, b)
        ]
        assert len(lignes[0]) == len(lignes[1]), nom
        # Chemins depuis la racine de chaque ligne (AD-7).
        return {d for x, y in zip(*lignes, strict=True) for d in _differences(x, y)}
    if relatif.endswith(".xlsx"):
        return _differences_cles(_cellules(a), _cellules(b))
    if relatif.endswith(".pptx"):
        return _differences_cles(_textes_slides(a), _textes_slides(b))
    raise AssertionError(f"format non comparé : {relatif}")  # pragma: no cover


def test_rejeu_composite_ne_differe_que_par_les_exclusions(
    racine: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    _simuler_composite(monkeypatch)
    caplog.set_level(logging.INFO, logger="usine.chaine.execution")
    for _ in range(2):
        code, erreur = _usine(monkeypatch, capsys)
        assert code == 0, erreur
    premier, second = sorted((racine / "sorties" / "passages" / DATE).iterdir())

    manifeste = json.loads((premier / FICHIERS_PASSAGE["manifeste"]).read_text(encoding="utf-8"))
    assert manifeste["palier"] == 1
    assert manifeste["source"] == "composite"
    statuts = {e["etape"]: e["statut"] for e in manifeste["etapes"]}
    for etape in composition.ETAPES_LLM:
        assert statuts[etape.value] == StatutEtape.sautee.value
    for p in (premier, second):
        assert (p / f"rapport-{DATE}.xlsx").is_file()
        assert (p / f"rapport-{DATE}.pptx").is_file()

    observees = {(nom, d) for nom in FICHIERS_PASSAGE for d in _comparer(nom, premier, second)}
    hors_exclusions = {(nom, d) for nom, d in observees if d not in EXCLUSIONS_REJEU.get(nom, ())}
    assert not hors_exclusions
    assert observees >= DIFFERENCES_ATTENDUES

    messages = [r.getMessage() for r in caplog.records if r.name == "usine.chaine.execution"]
    assert any("début" in m for m in messages)
    for etape in Etape:
        assert any(f"Étape « {etape.value} »" in m for m in messages), etape
    assert sum("publié" in m for m in messages) == 2


class _Panne:
    def executer(self, ctx: ContextePassage[Consultation]) -> ResultatEtape:
        raise RuntimeError("panne du rendu")


def test_echec_du_rendu_sans_rapport_publie(
    racine: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    construire = composition.construire_rapport

    def construire_en_panne(r: Path, *args: Any, **kwargs: Any) -> composition.Rapport:
        rapport = construire(r, *args, **kwargs)
        return replace(rapport, registre={**rapport.registre, Etape.rendu: _Panne()})

    monkeypatch.setattr(composition, "construire_rapport", construire_en_panne)
    code, erreur = _usine(monkeypatch, capsys)
    assert code == 1
    assert "rendu" in erreur
    assert not list((racine / "sorties" / "passages").rglob("rapport-*"))
    assert len(list((racine / "sorties" / "echecs").iterdir())) == 1
