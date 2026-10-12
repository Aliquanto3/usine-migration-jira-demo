"""`executer_passage` : Passage complet ou rien, progression, manifeste (Story 0.5)."""

import json
import re
from collections.abc import Callable
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation
from pydantic import BaseModel

from tests.doublures.fournisseur_factice import AppelEnregistre, FournisseurFactice
from tests.doublures.pseudonymiseur_factice import PseudonymiseurFactice
from tests.doublures.source_en_memoire import SourceEnMemoire
from tests.unit.chaine.environnement import (
    DATE_ARRETE,
    INSTANT,
    contenu_previsions,
    horloge_figee,
    horloge_pas_a_pas,
    preparer_racine,
    sans_exclusions,
)
from tests.unit.domaine.constructeur import Flux
from usine import composition
from usine.adaptateurs.journal import SourceJournal
from usine.adaptateurs.llm import REGISTRE
from usine.adaptateurs.presidio import PseudonymiseurPresidio
from usine.adaptateurs.pseudonymisation import MOTIFS_PAR_DEFAUT, PseudonymiseurHmac
from usine.chaine.ecrivain import EcrivainPassage
from usine.chaine.etapes import EtapeInstantanes, EtapeMeteo, EtapePrevision
from usine.chaine.execution import executer_passage
from usine.composition import PassageLance, Rapport, construire_rapport, lancer_rapport
from usine.consultation.passage import ConsultationPassage
from usine.contrats.calendrier import jours_ouvres
from usine.contrats.config import GrilleMeteo
from usine.contrats.erreurs import UsineErreur
from usine.contrats.llm import (
    ParametresAppel,
    ReponseClassification,
    ReponseExplication,
    ReponseNarration,
    ReponseStructuree,
    StatutControle,
    Tache,
)
from usine.contrats.meteo import CodeIndicateur, Couleur, MeteoEntite
from usine.contrats.passage import (
    CHEMIN_JOURNAL,
    FICHIERS_PASSAGE,
    MOTIF_PASSAGE_ID,
    ContextePassage,
    Etape,
    EvenementProgression,
    ManifestePassage,
    ResultatEtape,
    StatutEtape,
    lire_enveloppe,
)
from usine.contrats.prevision import ContenuPrevisions
from usine.contrats.tickets import CauseRacine, TypeTicket, rang_cle
from usine.contrats.trace import EntreeTrace, TypeTrace
from usine.contrats.versions import CONTRATS_VERSION, version_prompt
from usine.domaine.kpi import kpi as calculer_kpi
from usine.domaine.kpi import tendances as calculer_tendances
from usine.domaine.meteo import indicateurs as calculer_indicateurs
from usine.domaine.meteo import meteo as calculer_meteo
from usine.domaine.prevision.baseline import prevoir_sortie
from usine.domaine.prevision.risque import risques_entree_rouge
from usine.domaine.prevision.serie import series_stock
from usine.ia.etape import EtapeExplications
from usine.ia.explication import TEXTE_PLAFOND, lire_prompt
from usine.ia.narration import lire_prompt as lire_prompt_narration
from usine.ports.consultation import Consultation
from usine.ports.etape import Etape as PortEtape

ETAPES_LLM = {Etape.classification, Etape.repartition_causes, Etape.explications, Etape.narration}


class EtapeEnEchec:
    def __init__(self, exc: Exception) -> None:
        self._exc = exc

    def executer(self, ctx: ContextePassage[Consultation]) -> ResultatEtape:
        raise self._exc


class EtapeHorsSchema:
    def executer(self, ctx: ContextePassage[Consultation]) -> ResultatEtape:
        return ResultatEtape(StatutEtape.reussie, {"indicateurs": pl.DataFrame({"x": [1]})})


@pytest.fixture
def rapport(tmp_path: Path) -> Rapport:
    return construire_rapport(preparer_racine(tmp_path))


def _executer(
    rapport: Rapport,
    *,
    remplacements: dict[Etape, object] | None = None,
    horloge: Callable[[], datetime] | None = None,
    date_arrete: date = DATE_ARRETE,
) -> tuple[ManifestePassage, EcrivainPassage]:
    horloge = horloge or horloge_pas_a_pas()
    registre = {**rapport.registre, **(remplacements or {})}
    ecrivain = EcrivainPassage(rapport.racine_sorties, horloge)
    manifeste = executer_passage(
        registre,  # type: ignore[arg-type]
        rapport.configuration,
        date_arrete,
        ecrivain,
        ConsultationPassage,
        horloge,
        lambda: "abc123",
        rapport.donnees,
    )
    return manifeste, ecrivain


def _progression(ecrivain: EcrivainPassage) -> list[EvenementProgression]:
    lignes = ecrivain.chemin_progression.read_text(encoding="utf-8").split("\n")
    return [EvenementProgression.model_validate_json(ligne) for ligne in lignes if ligne]


def test_passage_publie(rapport: Rapport) -> None:
    manifeste, ecrivain = _executer(rapport)
    assert MOTIF_PASSAGE_ID.fullmatch(manifeste.passage_id)
    assert manifeste.passage_id == "20261012T073000Z-abc123"
    repertoire = rapport.racine_sorties / "passages" / "2026-10-09" / manifeste.passage_id
    assert ecrivain.repertoire_publie == repertoire
    assert sorted(p.name for p in repertoire.iterdir()) == [
        "indicateurs.parquet",
        "instantanes.parquet",
        "kpi.parquet",
        "manifeste.json",
        "meteo.json",
        "previsions.json",
        "rapport-2026-10-09.pptx",
        "rapport-2026-10-09.xlsx",
        "tendances.parquet",
        "tickets_etats.parquet",
        "tickets_textes.parquet",
        "trace",
    ]
    assert sorted(p.name for p in (repertoire / "trace").iterdir()) == ["previsions.jsonl"]
    assert not list((rapport.racine_sorties / "passages").glob(".en-cours-*"))
    relu = ManifestePassage.model_validate_json((repertoire / "manifeste.json").read_bytes())
    assert relu == manifeste


def test_manifeste(rapport: Rapport) -> None:
    manifeste, ecrivain = _executer(rapport)
    assert [e.etape for e in manifeste.etapes] == list(Etape)
    for execution in manifeste.etapes:
        attendu = StatutEtape.sautee if execution.etape in ETAPES_LLM else StatutEtape.reussie
        assert execution.statut is attendu
        assert execution.horodatage_fin > execution.horodatage_debut
    assert manifeste.palier == 1
    assert manifeste.version_grille == rapport.donnees.version_grille
    assert manifeste.version_donnees == rapport.donnees.version_donnees
    assert manifeste.source == "journal"
    assert ecrivain.repertoire_publie is not None
    etats = pl.read_parquet(ecrivain.repertoire_publie / "tickets_etats.parquet")
    nb_attendu = etats.filter(pl.col("date") == DATE_ARRETE).height
    assert manifeste.nb_tickets == nb_attendu > 0
    assert manifeste.horodatage_debut < manifeste.horodatage_fin
    assert manifeste.horodatage_extraction is None
    assert manifeste.nb_alignements == 0


class EtapeNotee:
    """Étape réussie sans sortie qui note son passage."""

    def __init__(self, journal: list[str], nom: str) -> None:
        self._journal = journal
        self._nom = nom

    def executer(self, ctx: ContextePassage[Consultation]) -> ResultatEtape:
        self._journal.append(self._nom)
        return ResultatEtape(StatutEtape.reussie)


class EtapeTracesLLM:
    """Étape réussie sans sortie qui rend des traces LLM : `(tâche, depuis_cache)`."""

    def __init__(self, appels: list[tuple[Tache, bool]]) -> None:
        self._appels = appels

    def executer(self, ctx: ContextePassage[Consultation]) -> ResultatEtape:
        traces = tuple(
            EntreeTrace(
                type_trace=TypeTrace.llm,
                horodatage=INSTANT,
                passage_id=ctx.passage_id,
                date_arrete=ctx.date_arrete,
                tache=tache,
                cible=f"DEMO-{rang}",
                fournisseur="factice",
                version_prompt="p" * 64,
                version_modele="factice:modele",
                version_grille="grille123456",
                version_donnees="d" * 64,
                entree_pseudonymisee="entrée",
                sortie="{}",
                statut_controle=StatutControle.conforme,
                depuis_cache=depuis_cache,
            )
            for rang, (tache, depuis_cache) in enumerate(self._appels, start=1)
        )
        return ResultatEtape(StatutEtape.reussie, traces=traces)


@pytest.mark.parametrize(
    ("appels", "attendu"),
    [
        ([], False),
        ([(Tache.classification, False), (Tache.classification, False)], True),
        ([(Tache.classification, False), (Tache.classification, True)], False),
        ([(Tache.classification, True)], False),
        ([(Tache.explication, False)], False),
    ],
)
def test_classification_initiale_du_manifeste(
    rapport: Rapport, appels: list[tuple[Tache, bool]], attendu: bool
) -> None:
    """Vrai si la classification a appelé le fournisseur sans aucun succès en cache."""
    manifeste, _ = _executer(rapport, remplacements={Etape.classification: EtapeTracesLLM(appels)})
    assert manifeste.classification_initiale is attendu


def test_extraction_et_alignements_recopies_au_manifeste(rapport: Rapport) -> None:
    """Constat 17 : `nb_alignements` et `horodatage_extraction` viennent de l'extraction."""
    ordre: list[str] = []

    def compter() -> int:
        ordre.append("compte")
        return 3

    extraction = INSTANT - timedelta(minutes=5)
    horloge = horloge_pas_a_pas()
    registre = {
        **rapport.registre,
        Etape.extraction: EtapeNotee(ordre, "extraction"),
        Etape.rendu: EtapeNotee(ordre, "rendu"),
    }
    manifeste = executer_passage(
        registre,
        rapport.configuration,
        DATE_ARRETE,
        EcrivainPassage(rapport.racine_sorties, horloge),
        ConsultationPassage,
        horloge,
        lambda: "abc123",
        rapport.donnees,
        horodatage_extraction=extraction,
        nb_alignements=compter,
    )
    assert ordre == ["extraction", "rendu", "compte"]
    assert manifeste.horodatage_extraction == extraction
    assert manifeste.nb_alignements == 3


def test_manifeste_complet_relu(rapport: Rapport) -> None:
    """AC 3 : chaque champ de traçabilité est présent dans `manifeste.json` relu."""
    manifeste, ecrivain = _executer(rapport)
    assert ecrivain.repertoire_publie is not None
    brut = json.loads((ecrivain.repertoire_publie / "manifeste.json").read_text(encoding="utf-8"))
    for champ in (
        "date_arrete",
        "source",
        "horodatage_extraction",
        "nb_tickets",
        "nb_alignements",
        "version_grille",
        "version_donnees",
        "etapes",
        "contrats_version",
    ):
        assert champ in brut, champ
    assert brut["date_arrete"] == "2026-10-09"
    assert brut["source"] == "journal"
    assert brut["horodatage_extraction"] is None  # aucune extraction Jira dans l'essai
    assert brut["nb_tickets"] == manifeste.nb_tickets > 0
    assert isinstance(brut["nb_alignements"], int)
    assert brut["nb_alignements"] == manifeste.nb_alignements >= 0
    assert brut["version_grille"] == rapport.donnees.version_grille
    assert brut["version_donnees"] == rapport.donnees.version_donnees
    assert brut["contrats_version"] == CONTRATS_VERSION
    assert [e["etape"] for e in brut["etapes"]] == [e.value for e in Etape]
    for execution in brut["etapes"]:
        assert execution["horodatage_debut"]
        assert execution["horodatage_fin"]
    consultation = ConsultationPassage(ecrivain.repertoire_publie)
    assert consultation.manifeste() == manifeste


def test_autoportance_kpi(rapport: Rapport) -> None:
    """AC 1 : KPI et Tendances publiés = sorties directes du domaine, chaque Jour ouvré."""
    manifeste, ecrivain = _executer(rapport)
    repertoire = ecrivain.repertoire_publie
    assert repertoire is not None
    racine = rapport.racine_sorties.parent
    ancrage = SourceJournal(racine / CHEMIN_JOURNAL).entete.ancrage_semaines
    programme = rapport.configuration.programme
    assert programme is not None
    entites = [entite.code for entite in programme.entites]
    etats = pl.read_parquet(repertoire / "tickets_etats.parquet")
    quotidiens = pl.read_parquet(repertoire / "kpi.parquet")
    hebdo = pl.read_parquet(repertoire / "tendances.parquet")
    dates = sorted(quotidiens.get_column("date").unique().to_list())
    assert dates == list(jours_ouvres(ancrage, DATE_ARRETE))
    directs = calculer_kpi(etats, entites)
    assert quotidiens.equals(directs)
    assert hebdo.equals(calculer_tendances(directs, DATE_ARRETE, ancrage))
    assert hebdo.get_column("date").max() == DATE_ARRETE
    consultation = ConsultationPassage(repertoire)
    assert consultation.kpi().equals(directs.filter(pl.col("date") == DATE_ARRETE))
    assert consultation.tendances().equals(hebdo)
    assert manifeste.date_arrete == DATE_ARRETE


def test_passage_avec_previsions(rapport: Rapport) -> None:
    """AC 2 : une `EtapePrevision` injectée publie `previsions.json` et ses traces."""
    contenu = contenu_previsions(DATE_ARRETE, rapport.donnees.version_grille)

    def prevoir(ctx: ContextePassage[Consultation]) -> ContenuPrevisions:
        # La prévision lit le Passage en cours : la Météo y est déjà.
        assert ctx.consultation.meteo(ctx.date_arrete)
        return contenu

    etape = EtapePrevision(prevoir, rapport.donnees.version_donnees, horloge_figee())
    manifeste, ecrivain = _executer(rapport, remplacements={Etape.prevision: etape})
    repertoire = ecrivain.repertoire_publie
    assert repertoire is not None
    assert (repertoire / "previsions.json").is_file()
    assert (repertoire / "trace" / "previsions.jsonl").is_file()
    assert not (repertoire / "trace" / "llm.jsonl").exists()
    consultation = ConsultationPassage(repertoire)
    assert consultation.previsions() == contenu.previsions
    assert consultation.risques() == contenu.risques
    traces = consultation.traces()
    assert [(t.type_trace, t.cible) for t in traces] == [
        (TypeTrace.prevision, "ENT-A"),
        (TypeTrace.prevision, "ENT-B"),
    ]
    assert {t.passage_id for t in traces} == {manifeste.passage_id}
    assert {t.horodatage for t in traces} == {INSTANT}
    assert {t.version_grille for t in traces} == {manifeste.version_grille}
    assert {t.version_donnees for t in traces} == {manifeste.version_donnees}


def test_previsions_du_passage(rapport: Rapport) -> None:
    """Story 5.3 : `previsions.json` = appel direct du domaine sur les séries du Passage.

    Une Prévision Baseline par Entité rouge et par Jour ouvré, les Risques des autres, dans
    l'ordre jour puis `meteo.json` ; une trace par élément.
    """
    manifeste, ecrivain = _executer(rapport)
    repertoire = ecrivain.repertoire_publie
    assert repertoire is not None
    grille = rapport.configuration.grille_meteo
    assert isinstance(grille, GrilleMeteo)
    etats = pl.read_parquet(repertoire / "tickets_etats.parquet")
    indicateurs = pl.read_parquet(repertoire / "indicateurs.parquet")
    meteos = lire_enveloppe(
        (repertoire / "meteo.json").read_bytes(), tuple[MeteoEntite, ...]
    ).contenu
    jours = sorted({m.date_arrete for m in meteos})
    assert jours == list(jours_ouvres(date(2026, 8, 17), DATE_ARRETE))
    series = series_stock(etats, ["ENT-A", "ENT-B"], jours)
    previsions = tuple(prevoir_sortie(series, m) for m in meteos if m.couleur is Couleur.rouge)
    risques = tuple(
        projete.risque
        for m in meteos
        if m.couleur is not Couleur.rouge
        for projete in risques_entree_rouge(indicateurs, m, grille).risques
    )
    publie = lire_enveloppe(
        (repertoire / "previsions.json").read_bytes(), ContenuPrevisions
    ).contenu
    assert publie == ContenuPrevisions(previsions=previsions, risques=risques)
    rouges = [(m.date_arrete, m.entite) for m in meteos if m.couleur is Couleur.rouge]
    assert rouges, "le jeu graine 42 doit avoir des Entités rouges"
    assert [(p.date_arrete, p.entite) for p in publie.previsions] == rouges
    # Graine 42, flux scripté (contrats 1.2) : ENT-B rouge à D, ENT-A sans Anomalie ouverte.
    assert {(m.date_arrete, m.entite) for m in meteos if m.couleur is Couleur.rouge} >= {
        (DATE_ARRETE, "ENT-B"),
    }
    consultation = ConsultationPassage(repertoire)
    traces = consultation.traces()
    assert len(traces) == len(publie.previsions) + len(publie.risques)
    assert {t.type_trace for t in traces} == {TypeTrace.prevision}
    assert [t.sortie for t in traces] == [
        e.model_dump_json() for e in (*publie.previsions, *publie.risques)
    ]
    assert {t.passage_id for t in traces} == {manifeste.passage_id}
    assert {t.version_donnees for t in traces} == {manifeste.version_donnees}


def test_progression(rapport: Rapport) -> None:
    _, ecrivain = _executer(rapport)
    lignes = _progression(ecrivain)
    attendu = []
    for etape in Etape:
        final = StatutEtape.sautee if etape in ETAPES_LLM else StatutEtape.reussie
        attendu += [(etape, StatutEtape.en_cours), (etape, final)]
    assert [(e.etape, e.statut) for e in lignes] == attendu
    assert {e.passage_id for e in lignes} == {ecrivain.passage_id}


def test_autoportance_meteo(rapport: Rapport) -> None:
    manifeste, ecrivain = _executer(rapport)
    assert ecrivain.repertoire_publie is not None
    texte = (ecrivain.repertoire_publie / "meteo.json").read_bytes()
    meteo = lire_enveloppe(texte, tuple[MeteoEntite, ...]).contenu
    jours = jours_ouvres(date(2026, 8, 17), DATE_ARRETE)
    attendu = [(jour, entite) for jour in jours for entite in ("ENT-A", "ENT-B")]
    assert [(m.date_arrete, m.entite) for m in meteo] == attendu
    assert {m.version_grille for m in meteo} == {manifeste.version_grille}
    indicateurs = pl.read_parquet(ecrivain.repertoire_publie / "indicateurs.parquet")
    # Grille livrée : IND-1 à IND-8 actifs, une ligne par indicateur, Entité et jour.
    codes = [code.value for code in CodeIndicateur]
    assert indicateurs.height == len(attendu) * len(codes)
    assert indicateurs.select("date", "entite").unique(maintain_order=True).rows() == attendu
    assert indicateurs["code"].cast(pl.Utf8).to_list() == codes * len(attendu)
    # Valeurs : identiques à un appel direct du domaine sur tout l'historique du Passage.
    grille = rapport.configuration.grille_meteo
    programme = rapport.configuration.programme
    assert isinstance(grille, GrilleMeteo)
    assert programme is not None
    racine = rapport.racine_sorties.parent
    ancrage = SourceJournal(racine / CHEMIN_JOURNAL).entete.ancrage_semaines
    etats = pl.read_parquet(ecrivain.repertoire_publie / "tickets_etats.parquet")
    directs = pl.concat(
        [
            calculer_indicateurs(
                etats.filter(pl.col("date") <= jour),
                jour,
                ["ENT-A", "ENT-B"],
                grille,
                programme=programme,
                ancrage=ancrage,
            )
            for jour in jours
        ],
        how="vertical",
    )
    colonnes = ["date", "entite", "statut", "motif", "numerateur", "denominateur"]
    comparees = [*colonnes, pl.col("code").cast(pl.Utf8), pl.col("couleur").cast(pl.Utf8)]
    assert indicateurs.select(comparees).equals(directs.select(comparees))
    # Météo publiée (Story 3.6) : chaque `MeteoEntite` de `meteo.json`, règles, valeurs et
    # tickets en cause compris, est celle que calcule directement le domaine.
    meteos_directes = tuple(
        m
        for jour in jours
        for m in calculer_meteo(
            etats.filter(pl.col("date") <= jour),
            jour,
            ["ENT-A", "ENT-B"],
            grille,
            manifeste.version_grille,
            programme=programme,
            ancrage=ancrage,
        )
    )
    assert meteo == meteos_directes
    assert any(m.regles for m in meteo)
    assert all(regle.tickets_en_cause for m in meteo for regle in m.regles)


@pytest.mark.parametrize(
    ("exc", "message"),
    [
        (RuntimeError("panne"), "RuntimeError : panne"),
        (UsineErreur("message d'origine"), "message d'origine"),
    ],
)
def test_echec_d_etape(rapport: Rapport, exc: Exception, message: str) -> None:
    with pytest.raises(UsineErreur, match="kpi") as erreur:
        _executer(rapport, remplacements={Etape.kpi: EtapeEnEchec(exc)})
    assert message in str(erreur.value)
    passage_id = "20261012T073000Z-abc123"
    echec = rapport.racine_sorties / "echecs" / passage_id
    assert (echec / "tickets_etats.parquet").exists()
    assert not (echec / "manifeste.json").exists()
    assert not (rapport.racine_sorties / "passages" / "2026-10-09").exists()
    assert not list((rapport.racine_sorties / "passages").glob(".en-cours-*"))
    progression = rapport.racine_sorties / "progression" / f"{passage_id}.jsonl"
    derniere = EvenementProgression.model_validate_json(
        progression.read_text(encoding="utf-8").splitlines()[-1]
    )
    assert (derniere.etape, derniere.statut) == (Etape.kpi, StatutEtape.echouee)
    assert message in derniere.message


def test_sortie_hors_schema(rapport: Rapport) -> None:
    with pytest.raises(UsineErreur, match=r"kpi.*SCHEMA_INDICATEUR"):
        _executer(rapport, remplacements={Etape.kpi: EtapeHorsSchema()})
    assert (rapport.racine_sorties / "echecs" / "20261012T073000Z-abc123").is_dir()


def test_registre_incomplet(rapport: Rapport) -> None:
    registre: dict[Etape, PortEtape] = {
        e: i for e, i in rapport.registre.items() if e is not Etape.rendu
    }
    with pytest.raises(UsineErreur, match="rendu"):
        _executer(replace(rapport, registre=registre))
    assert not rapport.racine_sorties.exists()


def test_horloge_figee(rapport: Rapport) -> None:
    manifeste, _ = _executer(rapport, horloge=horloge_figee())
    assert all(e.duree.total_seconds() == 0 for e in manifeste.etapes)


def test_jours_sans_ticket_verts(rapport: Rapport) -> None:
    """Premiers jours et Date d'arrêté sans ticket créé : Météo verte, sans erreur."""
    flux = Flux()
    flux.anomalie("DEMO-1", datetime(2026, 10, 7, 9, tzinfo=UTC))  # mercredi
    flux.transition("DEMO-1", datetime(2026, 10, 8, 9, tzinfo=UTC), "nouvelle", "rejetee")
    grille = rapport.configuration.grille_meteo
    programme = rapport.configuration.programme
    assert isinstance(grille, GrilleMeteo)
    assert programme is not None
    premier = date(2026, 10, 5)
    remplacements: dict[Etape, object] = {
        Etape.instantanes: EtapeInstantanes(SourceEnMemoire(flux.evenements), premier),
        Etape.meteo: EtapeMeteo(grille, "abcdef012345", ["ENT-A", "ENT-B"], premier, programme),
    }
    _, ecrivain = _executer(rapport, remplacements=remplacements)
    assert ecrivain.repertoire_publie is not None
    texte = (ecrivain.repertoire_publie / "meteo.json").read_bytes()
    couleurs = {
        (m.date_arrete.day, m.entite): m.couleur
        for m in lire_enveloppe(texte, tuple[MeteoEntite, ...]).contenu
    }
    assert couleurs[(5, "ENT-A")] is Couleur.vert
    assert couleurs[(7, "ENT-A")] is Couleur.orange
    assert couleurs[(9, "ENT-A")] is Couleur.vert
    assert all(couleurs[(jour, "ENT-B")] is Couleur.vert for jour in range(5, 10))
    # Date d'arrêté antérieure à toute création : aucune ligne d'état, tout est vert.
    plus_tard = horloge_pas_a_pas(datetime(2026, 10, 12, 8, tzinfo=UTC))
    manifeste, ecrivain = _executer(
        rapport, remplacements=remplacements, horloge=plus_tard, date_arrete=date(2026, 10, 6)
    )
    assert manifeste.nb_tickets == 0
    assert ecrivain.repertoire_publie is not None
    texte = (ecrivain.repertoire_publie / "meteo.json").read_bytes()
    meteo = lire_enveloppe(texte, tuple[MeteoEntite, ...]).contenu
    assert len(meteo) == 4
    assert {m.couleur for m in meteo} == {Couleur.vert}


def test_date_anterieure_a_l_historique(rapport: Rapport) -> None:
    with pytest.raises(UsineErreur, match=r"instantanes.*2026-08-14"):
        _executer(rapport, date_arrete=date(2026, 8, 14))


ID = "20261012T073000Z-abc123"


def _lignes_progression(rapport: Rapport) -> list[EvenementProgression]:
    chemin = rapport.racine_sorties / "progression" / f"{ID}.jsonl"
    lignes = chemin.read_text(encoding="utf-8").splitlines()
    return [EvenementProgression.model_validate_json(ligne) for ligne in lignes]


def test_interruption_nettoyee_et_propagee(rapport: Rapport) -> None:
    with pytest.raises(KeyboardInterrupt):
        _executer(rapport, remplacements={Etape.kpi: EtapeEnEchec(KeyboardInterrupt())})  # type: ignore[arg-type]
    assert (rapport.racine_sorties / "echecs" / ID).is_dir()
    assert not list((rapport.racine_sorties / "passages").glob(".en-cours-*"))
    derniere = _lignes_progression(rapport)[-1]
    assert (derniere.etape, derniere.statut) == (Etape.kpi, StatutEtape.echouee)


def test_echec_de_publication(rapport: Rapport, monkeypatch: pytest.MonkeyPatch) -> None:
    def publier_en_panne(self: EcrivainPassage, manifeste: ManifestePassage) -> Path:
        raise OSError("disque plein")

    monkeypatch.setattr(EcrivainPassage, "publier", publier_en_panne)
    with pytest.raises(UsineErreur, match="publication du Passage") as erreur:
        _executer(rapport)
    assert "disque plein" in str(erreur.value)
    derniere = _lignes_progression(rapport)[-1]
    assert (derniere.etape, derniere.statut) == (Etape.rendu, StatutEtape.reussie)
    assert (rapport.racine_sorties / "echecs" / ID).is_dir()
    assert not (rapport.racine_sorties / "passages" / "2026-10-09").exists()


def test_abandon_en_panne_garde_l_erreur_d_etape(
    rapport: Rapport, monkeypatch: pytest.MonkeyPatch
) -> None:
    def abandonner_en_panne(self: EcrivainPassage) -> Path | None:
        raise UsineErreur("déplacement impossible")

    monkeypatch.setattr(EcrivainPassage, "abandonner", abandonner_en_panne)
    with pytest.raises(UsineErreur, match="kpi") as erreur:
        _executer(rapport, remplacements={Etape.kpi: EtapeEnEchec(RuntimeError("panne"))})
    assert "panne" in str(erreur.value)
    assert "déplacement impossible" not in str(erreur.value)


def _passer_en_palier_2(racine: Path) -> None:
    chemin = racine / "config" / "chaine.yaml"
    texte = chemin.read_text(encoding="utf-8")
    assert "palier: 1" in texte
    chemin.write_text(texte.replace("palier: 1", "palier: 2"), encoding="utf-8")


def test_palier_2_explications(tmp_path: Path) -> None:
    # Météo de référence (Palier 1, même Journal) : elle fixe les réponses scriptées.
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    reponses = _reponses(a_expliquer)

    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    fournisseur = FournisseurPalier2(reponses)
    rapport = construire_rapport(
        racine, fournisseur, PseudonymiseurFactice(b"cle-fictive", {}), horloge_figee()
    )
    manifeste, ecrivain = _executer(rapport)
    assert manifeste.palier == 2
    statuts = {e.etape: e.statut for e in manifeste.etapes}
    assert statuts[Etape.explications] is StatutEtape.reussie
    assert statuts[Etape.classification] is StatutEtape.reussie
    assert statuts[Etape.repartition_causes] is StatutEtape.reussie
    assert statuts[Etape.narration] is StatutEtape.reussie
    repertoire = ecrivain.repertoire_publie
    assert repertoire is not None
    # Story 6.5 : répartition publiée, `nombre` entier, parts de somme 1 par Entité.
    repartition = ConsultationPassage(repertoire).repartition_causes()
    assert not repartition.is_empty()
    assert repartition["nombre"].dtype == pl.Int64()
    assert set(repartition["cause_racine"].cast(pl.Utf8)) == {CAUSE_SCRIPTEE.value}
    sommes = repartition.group_by("entite").agg(pl.col("part").sum())["part"]
    assert all(abs(somme - 1.0) < 1e-9 for somme in sommes)
    assert (repertoire / "llm" / "explications.json").is_file()
    assert (repertoire / "trace" / "llm.jsonl").is_file()
    consultation = ConsultationPassage(repertoire)
    explications = consultation.explications()
    assert [e.entite for e in explications] == [m.entite for m in a_expliquer]
    assert all(e.statut_controle is StatutControle.conforme for e in explications)
    traces = [t for t in consultation.traces() if t.tache is Tache.explication]
    assert len(traces) == len(a_expliquer) == len(_appels(fournisseur, Tache.explication))
    assert {t.passage_id for t in traces} == {manifeste.passage_id}
    # Story 6.7 : une narration par Entité ayant une Prévision ou des Risques.
    assert (repertoire / "llm" / "narrations.json").is_file()
    narrations = consultation.narrations()
    assert [n.entite for n in narrations] == list(a_expliquer.a_narrer)
    assert all(n.statut_controle is StatutControle.conforme for n in narrations)
    assert len(_appels(fournisseur, Tache.narration)) == len(a_expliquer.a_narrer)
    assert {t.version_donnees for t in traces} == {rapport.donnees.version_donnees}
    assert (racine / "donnees" / "cache_llm.sqlite").is_file()


def test_palier_2_brouillon_dans_excel_et_powerpoint(tmp_path: Path) -> None:
    """Story 6.6 : Brouillon marqué dans l'onglet Commentaires et sur la slide de l'Entité."""
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    brouillons = {
        m.entite: f"Relancer l'équipe sur {m.regles[0].tickets_en_cause[0]}." for m in a_expliquer
    }
    reponses = [
        ReponseExplication(
            texte=f"{m.regles[0].tickets_en_cause[0]} est en cause.",
            brouillon_commentaire=brouillons[m.entite],
        )
        for m in a_expliquer
    ] + _reponses_narration(a_expliquer)
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    fournisseur = FournisseurPalier2(reponses)
    rapport = construire_rapport(
        racine, fournisseur, PseudonymiseurFactice(b"cle-fictive", {}), horloge_figee()
    )
    _, ecrivain = _executer(rapport)
    repertoire = ecrivain.repertoire_publie
    assert repertoire is not None

    # Prompt : KPI de l'Entité lus, répartition des Causes racines avec compte et part (6.5).
    for appel in _appels(fournisseur, Tache.explication):
        prompt = str(appel.entree)
        assert "### KPI de l'Entité\n\n- " in prompt
        assert "### KPI de l'Entité\n\n- aucun" not in prompt
        assert f"### Causes racines proposées\n\n- {CAUSE_SCRIPTEE.value} : " in prompt
        assert "(100 %)" in prompt
        assert "- non disponible" not in prompt

    attendus = {entite: f"brouillon IA, à valider : {b}" for entite, b in brouillons.items()}
    classeur = load_workbook(repertoire / "rapport-2026-10-09.xlsx")
    lignes = [
        [cellule for cellule in ligne if cellule is not None]
        for ligne in classeur["Commentaires"].iter_rows(values_only=True)
    ]
    for entite, attendu in attendus.items():
        assert any(ligne[:1] == [entite] and attendu in ligne for ligne in lignes), entite

    zones: dict[str, list[str]] = {}
    for slide in Presentation(str(repertoire / "rapport-2026-10-09.pptx")).slides:
        titre = slide.shapes.title
        if titre is None or " — " not in titre.text_frame.text:
            continue
        zones[titre.text_frame.text.split(" — ")[0]] = [
            forme.text_frame.text for forme in slide.shapes if forme.has_text_frame
        ]
    for entite, attendu in attendus.items():
        assert any(attendu in zone for zone in zones[entite]), entite


def test_palier_2_sans_fournisseur(tmp_path: Path) -> None:
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    (racine / "config" / "llm.yaml").unlink(missing_ok=True)
    with pytest.raises(UsineErreur, match="aucun fournisseur LLM"):
        construire_rapport(racine)
    assert not (racine / "sorties").exists()
    assert not (racine / "donnees" / "cache_llm.sqlite").exists()


def test_palier_2_fournisseur_de_llm_yaml_sans_cle(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """`config/llm.yaml` présent, `ANTHROPIC_API_KEY` absente : refus avant tout Passage."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    (racine / "config" / "llm.yaml").write_text(
        "fournisseur: anthropic\nmodele: claude-sonnet-5-5\nplafond_cout_passage_eur: 1\n",
        encoding="utf-8",
    )
    with pytest.raises(UsineErreur, match="ANTHROPIC_API_KEY"):
        construire_rapport(racine)
    assert not (racine / "sorties").exists()


def test_palier_2_sans_cle_de_pseudonymisation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("USINE_CLE_PSEUDO", raising=False)
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    with pytest.raises(UsineErreur, match="USINE_CLE_PSEUDO"):
        construire_rapport(racine, FournisseurFactice())
    assert not (racine / "sorties").exists()
    assert not (racine / "donnees" / "cache_llm.sqlite").exists()
    assert not (racine / "donnees" / "pseudonymes.sqlite").exists()


NUMERO_COMPTE = "FR76 3000 6000 0112 3456 7890 189"
"""Numéro de compte fictif (IBAN de démonstration)."""


class _Reference(list[MeteoEntite]):
    """Météos rouges ou orange d'un Passage de référence, et Entités dont la Prévision est
    narrée (Story 6.7), dans l'ordre des appels au fournisseur."""

    a_narrer: tuple[str, ...] = ()


def _meteos_a_expliquer(racine: Path) -> _Reference:
    """Météos rouges ou orange d'un Passage Palier 1 sur le Journal de `racine`."""
    reference, _ = _executer(construire_rapport(racine))
    repertoire = racine / "sorties" / "passages" / "2026-10-09" / reference.passage_id
    consultation = ConsultationPassage(repertoire)
    a_expliquer = _Reference(
        m for m in consultation.meteo(DATE_ARRETE) if m.couleur in (Couleur.rouge, Couleur.orange)
    )
    assert a_expliquer, "le jeu de graine 42 doit avoir une Entité rouge ou orange"
    a_expliquer.a_narrer = tuple(
        sorted(
            {p.entite for p in consultation.previsions(DATE_ARRETE)}
            | {r.entite for r in consultation.risques(DATE_ARRETE)}
        )
    )
    assert a_expliquer.a_narrer, "le jeu de graine 42 doit avoir une Prévision ou un Risque"
    return a_expliquer


def _ajouter_compte_a_la_description(
    racine: Path, cle: str, ajout: str = f"Compte {NUMERO_COMPTE} concerné."
) -> None:
    """Ajoute `ajout` (numéro de compte fictif) à chaque description de `cle` dans le Journal."""
    chemin = racine / "donnees" / "journal.jsonl"
    lignes = chemin.read_text(encoding="utf-8").splitlines()
    modifiees = 0
    for rang, ligne in enumerate(lignes[1:], start=1):
        evenement = json.loads(ligne)
        if evenement["cle_ticket"] == cle and evenement["champ"] == "description":
            evenement["vers"] = f"{evenement['vers']} {ajout}"
            lignes[rang] = json.dumps(evenement, ensure_ascii=False, separators=(",", ":"))
            modifiees += 1
    assert modifiees, f"aucune description pour {cle}"
    chemin.write_text("\n".join(lignes) + "\n", encoding="utf-8")


def _reponses_narration(reference: _Reference) -> list[BaseModel]:
    """Une narration sans chiffre (conforme) par Entité à narrer, après les explications."""
    return [ReponseNarration(texte="La Prévision suit sa tendance.") for _ in reference.a_narrer]


def _reponses(a_expliquer: _Reference) -> list[BaseModel]:
    explications: list[BaseModel] = [
        ReponseExplication(texte=f"{m.regles[0].tickets_en_cause[0]} est en cause.")
        for m in a_expliquer
    ]
    return explications + _reponses_narration(a_expliquer)


CAUSE_SCRIPTEE = CauseRacine.parametrage
"""Cause racine que `FournisseurPalier2` propose pour chaque Anomalie."""


class FournisseurPalier2(FournisseurFactice):
    """Fournisseur factice qui répond seul à chaque classification ; les explications suivent
    les réponses scriptées, dans l'ordre."""

    def generer_structure[T: BaseModel](
        self, prompt: str, schema: type[T], parametres: ParametresAppel
    ) -> ReponseStructuree[T]:
        if schema is ReponseClassification:
            reponse = ReponseClassification(
                cause_racine=CAUSE_SCRIPTEE, justification="Paramètre manquant."
            )
            self._structurees.appendleft(reponse)
        return super().generer_structure(prompt, schema, parametres)


def _appels(fournisseur: FournisseurFactice, tache: Tache) -> list[AppelEnregistre]:
    """Appels structurés reçus par `fournisseur` pour la tâche `tache`."""
    return [a for a in fournisseur.appels if a.parametres.tache is tache]


def test_palier_2_par_lancer_rapport(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    instant = datetime(2026, 10, 12, 6, 45, tzinfo=UTC)
    fournisseur = FournisseurPalier2(_reponses(a_expliquer))
    lance = lancer_rapport(
        racine, DATE_ARRETE, horloge_figee(instant), lambda: "abc123", fournisseur
    )
    assert lance.manifeste.palier == 2
    traces = [
        t for t in ConsultationPassage(lance.repertoire).traces() if t.tache is Tache.explication
    ]
    assert len(traces) == len(a_expliquer) == len(_appels(fournisseur, Tache.explication))
    assert {t.horodatage for t in traces} == {instant}
    assert {t.passage_id for t in traces} == {lance.manifeste.passage_id}
    assert (racine / "donnees" / "pseudonymes.sqlite").is_file()  # `PseudonymiseurHmac`


@pytest.mark.parametrize("pseudonymiseur", ["hmac", "factice"])
def test_deux_passages_palier_2_cache_et_pseudonymisation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, pseudonymiseur: str
) -> None:
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    cle_en_cause = a_expliquer[0].regles[0].tickets_en_cause[0]
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    _ajouter_compte_a_la_description(racine, cle_en_cause)

    def lancer(fournisseur: FournisseurFactice) -> PassageLance:
        injecte = (
            PseudonymiseurFactice(b"cle-de-test-fictive", {"COMPTE": MOTIFS_PAR_DEFAUT["COMPTE"]})
            if pseudonymiseur == "factice"
            else None
        )
        return lancer_rapport(
            racine,
            DATE_ARRETE,
            horloge_pas_a_pas(),
            fournisseur=fournisseur,
            pseudonymiseur=injecte,
        )

    premier_fournisseur = FournisseurPalier2(_reponses(a_expliquer))
    premier = lancer(premier_fournisseur)
    second_fournisseur = FournisseurFactice()  # aucune réponse : tout doit venir du cache
    second = lancer(second_fournisseur)
    assert second.repertoire != premier.repertoire
    assert second_fournisseur.appels == []

    # Le numéro de compte a atteint le prompt, pseudonymisé.
    prompts = [str(appel.entree) for appel in premier_fournisseur.appels]
    assert any(re.search(r"\bCOMPTE_[0-9a-f]{6}\b", prompt) for prompt in prompts)
    assert all(NUMERO_COMPTE not in prompt for prompt in prompts)

    gabarits = {
        Tache.explication: version_prompt(lire_prompt()),
        Tache.narration: version_prompt(lire_prompt_narration()),
    }
    for lance, depuis_cache in ((premier, False), (second, True)):
        lignes = (lance.repertoire / "trace" / "llm.jsonl").read_text(encoding="utf-8")
        toutes = [EntreeTrace.model_validate_json(ligne) for ligne in lignes.splitlines()]
        assert {t.depuis_cache for t in toutes} == {depuis_cache}
        traces = [t for t in toutes if t.tache is Tache.explication]
        assert len(traces) == len(a_expliquer)
        assert {t.depuis_cache for t in traces} == {depuis_cache}
        meteos = {m.entite: m for m in ConsultationPassage(lance.repertoire).meteo(DATE_ARRETE)}
        for trace in traces:
            assert trace.entite is not None
            assert trace.version_grille == meteos[trace.entite].version_grille
            assert trace.tache is not None
            assert trace.version_prompt == gabarits[trace.tache]
        for fichier in lance.repertoire.rglob("*"):
            if not fichier.is_file():
                continue
            if fichier.name == "tickets_textes.parquet":
                # Données sources du Passage (AD-23), locales et fictives : le numéro y est.
                textes = pl.read_parquet(fichier)
                assert textes["description"].str.contains(NUMERO_COMPTE, literal=True).any()
                continue
            assert NUMERO_COMPTE.encode("utf-8") not in fichier.read_bytes(), fichier.name


def test_rejeu_palier_2_exclusions_par_fichier(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """AD-7 : deux Passages Palier 2 (le second servi par le cache) ne diffèrent que par les
    champs exclus de `explications` et de `trace_llm`."""
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    premier = lancer_rapport(
        racine,
        DATE_ARRETE,
        horloge_pas_a_pas(),
        fournisseur=FournisseurPalier2(_reponses(a_expliquer)),
    )
    second_fournisseur = FournisseurFactice()  # aucune réponse : tout vient du cache
    second = lancer_rapport(
        racine,
        DATE_ARRETE,
        horloge_pas_a_pas(INSTANT + timedelta(hours=1)),
        fournisseur=second_fournisseur,
    )
    assert second_fournisseur.appels == []

    def lire(lance: PassageLance, nom: str) -> object:
        texte = (lance.repertoire / FICHIERS_PASSAGE[nom]).read_text(encoding="utf-8")
        if nom.startswith("trace_"):
            return [json.loads(ligne) for ligne in texte.splitlines() if ligne.strip()]
        return json.loads(texte)

    for nom in ("classification", "explications", "trace_llm"):
        a, b = lire(premier, nom), lire(second, nom)
        assert a != b, nom  # passage_id, horodatage, depuis_cache… diffèrent bien
        if nom == "trace_llm":
            assert isinstance(a, list) and isinstance(b, list)
            assert [sans_exclusions(nom, x) for x in a] == [sans_exclusions(nom, y) for y in b]
        else:
            assert sans_exclusions(nom, a) == sans_exclusions(nom, b), nom


def test_palier_2_classification_initiale_puis_incrementale(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """FR-35, AD-8 : le premier Passage Palier 2 classe tout le stock (`classification_initiale`)
    et publie `llm/classification.json` ; le suivant, à textes identiques, n'appelle pas."""
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    fournisseur = FournisseurPalier2(_reponses(a_expliquer))
    premier = lancer_rapport(racine, DATE_ARRETE, horloge_pas_a_pas(), fournisseur=fournisseur)
    assert premier.manifeste.classification_initiale is True
    assert (premier.repertoire / FICHIERS_PASSAGE["classification"]).is_file()
    consultation = ConsultationPassage(premier.repertoire)
    textes = consultation.textes_tickets()
    anomalies = sorted(
        textes.filter(pl.col("type_ticket") == TypeTicket.anomalie.value)["cle_ticket"],
        key=rang_cle,
    )
    assert anomalies
    classification = consultation.classification()
    assert [c.cle_ticket for c in classification] == anomalies
    assert {c.cause_racine for c in classification} == {CAUSE_SCRIPTEE}
    assert {c.statut_controle for c in classification} == {StatutControle.conforme}
    assert len(_appels(fournisseur, Tache.classification)) == len(anomalies)
    traces = [t for t in consultation.traces() if t.tache is Tache.classification]
    assert [t.cible for t in traces] == anomalies
    assert not any(t.depuis_cache for t in traces)

    second_fournisseur = FournisseurFactice()  # aucune réponse : tout vient du cache
    second = lancer_rapport(
        racine,
        DATE_ARRETE,
        horloge_pas_a_pas(INSTANT + timedelta(hours=1)),
        fournisseur=second_fournisseur,
    )
    assert second_fournisseur.appels == []
    assert second.manifeste.classification_initiale is False
    assert ConsultationPassage(second.repertoire).classification() == classification


def test_palier_2_cle_de_pseudonymisation_blanche(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("USINE_CLE_PSEUDO", "   ")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    with pytest.raises(UsineErreur, match="USINE_CLE_PSEUDO"):
        construire_rapport(racine, FournisseurFactice())
    assert not (racine / "donnees" / "pseudonymes.sqlite").exists()


def test_palier_2_pseudonymisation_sans_reconnaisseur(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    (racine / "config" / "pseudonymisation.yaml").write_text(
        "langue: fr\nreconnaisseurs: []\n", encoding="utf-8"
    )
    with pytest.raises(UsineErreur, match=r"pseudonymisation\.yaml.*reconnaisseurs"):
        construire_rapport(racine, FournisseurFactice())
    assert not (racine / "sorties").exists()
    assert not (racine / "donnees" / "pseudonymes.sqlite").exists()


def test_palier_2_deux_reconnaisseurs_du_meme_type(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    autre_compte = "cpt12345678"
    for meteo_entite in a_expliquer:
        _ajouter_compte_a_la_description(
            racine,
            meteo_entite.regles[0].tickets_en_cause[0],
            f"Comptes {NUMERO_COMPTE} et {autre_compte}.",
        )
    (racine / "config" / "pseudonymisation.yaml").write_text(
        "reconnaisseurs:\n"
        "  - type_jeton: COMPTE\n"
        f"    motif: '{MOTIFS_PAR_DEFAUT['COMPTE']}'\n"
        "  - type_jeton: COMPTE\n"
        "    motif: '\\bcpt\\d{8}\\b'\n",
        encoding="utf-8",
    )
    fournisseur = FournisseurPalier2(_reponses(a_expliquer))
    lancer_rapport(racine, DATE_ARRETE, horloge_figee(), fournisseur=fournisseur)
    entrees = [str(appel.entree) for appel in _appels(fournisseur, Tache.explication)]
    assert len(entrees) == len(a_expliquer)
    for entree in entrees:
        assert NUMERO_COMPTE not in entree
        assert autre_compte not in entree
        assert len(set(re.findall(r"\bCOMPTE_[0-9a-f]{6}\b", entree))) >= 2


def _ecrire_plafond(racine: Path, plafond: str) -> None:
    chemin = racine / "config" / "llm.yaml"
    texte = chemin.read_text(encoding="utf-8")
    assert "plafond_cout_passage_eur: 2.00" in texte
    chemin.write_text(
        texte.replace("plafond_cout_passage_eur: 2.00", f"plafond_cout_passage_eur: {plafond}"),
        encoding="utf-8",
    )


def test_palier_2_cout_llm_du_manifeste(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """FR-34, NFR-9 : chaque trace LLM porte son coût ; le manifeste porte leur somme."""
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    version = "anthropic:claude-sonnet-5-5"  # au registre des prix : coût non nul
    fournisseur = FournisseurPalier2(_reponses(a_expliquer), version_modele=version)
    lance = lancer_rapport(racine, DATE_ARRETE, horloge_pas_a_pas(), fournisseur=fournisseur)
    traces = [
        t for t in ConsultationPassage(lance.repertoire).traces() if t.type_trace is TypeTrace.llm
    ]
    assert traces and all(t.cout_estime_eur > 0 for t in traces)
    assert lance.manifeste.cout_llm_eur == sum((t.cout_estime_eur for t in traces), Decimal(0))
    tarif = REGISTRE[version]
    attendu = (
        tarif.prix_entree_eur_mtok * traces[0].jetons_entree
        + tarif.prix_sortie_eur_mtok * traces[0].jetons_sortie
    ) / 1_000_000
    assert traces[0].cout_estime_eur == attendu

    # Rejeu : tout vient du cache, coût nul.
    rejeu = lancer_rapport(
        racine,
        DATE_ARRETE,
        horloge_pas_a_pas(INSTANT + timedelta(hours=1)),
        fournisseur=FournisseurFactice(version_modele=version),
    )
    assert rejeu.manifeste.cout_llm_eur == 0


def test_palier_2_plafond_nul_explications_degradees(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """`plafond_cout_passage_eur: 0` : aucun appel, libellé de plafond, Passage publié."""
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    _ecrire_plafond(racine, "0")
    fournisseur = FournisseurFactice()
    lance = lancer_rapport(racine, DATE_ARRETE, horloge_figee(), fournisseur=fournisseur)
    assert fournisseur.appels == []
    explications = ConsultationPassage(lance.repertoire).explications()
    assert explications
    assert {e.texte for e in explications} == {TEXTE_PLAFOND}
    assert {e.statut_controle for e in explications} == {StatutControle.rejete}
    # Story 6.7 : la Prévision brute est publiée, sans échec du Passage.
    narrations = ConsultationPassage(lance.repertoire).narrations()
    assert narrations
    assert all(n.texte.startswith("Prévision brute.") for n in narrations)
    assert {n.statut_controle for n in narrations} == {StatutControle.rejete}
    statuts = {e.etape: e.statut for e in lance.manifeste.etapes}
    assert statuts[Etape.explications] is StatutEtape.reussie
    assert statuts[Etape.narration] is StatutEtape.reussie
    assert lance.manifeste.cout_llm_eur == 0


def test_palier_2_fournisseur_en_echec_nomme_l_etape(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """H-31 : une erreur du fournisseur fait échouer le Passage en nommant l'étape."""
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    with pytest.raises(UsineErreur, match="étape « classification »"):
        lancer_rapport(racine, DATE_ARRETE, horloge_figee(), fournisseur=FournisseurFactice())
    publies = racine / "sorties" / "passages" / DATE_ARRETE.isoformat()
    assert not publies.exists() or not any(publies.iterdir())


class _FournisseurFermable(FournisseurPalier2):
    """Fournisseur factice qui compte ses fermetures (adaptateur construit par la composition)."""

    def __init__(self, reponses: list[BaseModel] | None = None) -> None:
        super().__init__(reponses or ())
        self.fermetures = 0

    def fermer(self) -> None:
        self.fermetures += 1


@pytest.mark.parametrize("reussite", [True, False])
def test_fournisseur_construit_ferme_en_fin_de_passage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, reussite: bool
) -> None:
    a_expliquer = _meteos_a_expliquer(preparer_racine(tmp_path / "p1"))
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path / "p2")
    _passer_en_palier_2(racine)
    construit = _FournisseurFermable(_reponses(a_expliquer) if reussite else None)
    monkeypatch.setattr(composition, "construire_fournisseur", lambda *_: construit)
    monkeypatch.setattr(composition, "construire_fournisseurs_par_tache", lambda *_: {})
    if reussite:
        lancer_rapport(racine, DATE_ARRETE, horloge_figee())
    else:
        with pytest.raises(UsineErreur, match="explications"):
            lancer_rapport(racine, DATE_ARRETE, horloge_figee())
    assert construit.fermetures == 1


def test_fournisseur_construit_ferme_si_la_construction_echoue(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv("USINE_CLE_PSEUDO", raising=False)
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    construit = _FournisseurFermable()
    par_tache = _FournisseurFermable()
    monkeypatch.setattr(composition, "construire_fournisseur", lambda *_: construit)
    monkeypatch.setattr(
        composition,
        "construire_fournisseurs_par_tache",
        lambda *_: {Tache.classification: par_tache},
    )
    with pytest.raises(UsineErreur, match="USINE_CLE_PSEUDO"):
        construire_rapport(racine)
    assert (construit.fermetures, par_tache.fermetures) == (1, 1)


def test_fournisseur_construit_ferme_si_un_fournisseur_par_tache_echoue(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Story 6.5 : un modèle par tâche hors registre ferme le fournisseur déjà construit."""
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "cle-de-test")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    llm = racine / "config" / "llm.yaml"
    llm.write_text(
        llm.read_text(encoding="utf-8").replace(
            "classification: claude-haiku-5-5", "classification: modele-inconnu"
        ),
        encoding="utf-8",
    )
    construit = _FournisseurFermable()
    monkeypatch.setattr(composition, "construire_fournisseur", lambda *_: construit)
    with pytest.raises(UsineErreur, match="modele-inconnu"):
        construire_rapport(racine)
    assert construit.fermetures == 1


def test_fournisseur_injecte_jamais_ferme(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    injecte = _FournisseurFermable()
    with pytest.raises(UsineErreur):
        lancer_rapport(racine, DATE_ARRETE, horloge_figee(), fournisseur=injecte)
    assert injecte.fermetures == 0


def _pseudonymiseur_du_rapport(rapport: Rapport) -> object:
    etape = rapport.registre[Etape.explications]
    assert isinstance(etape, EtapeExplications)
    return etape._passerelle._pseudonymiseur


@pytest.mark.parametrize(
    ("moteur", "attendu"),
    [("regex", PseudonymiseurHmac), ("presidio", PseudonymiseurPresidio), (None, None)],
)
def test_moteur_de_pseudonymisation_choisi_par_la_configuration(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, moteur: str | None, attendu: type | None
) -> None:
    """AD-15 : `moteur` de `config/pseudonymisation.yaml` ; absent, `presidio`."""
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    chemin = racine / "config" / "pseudonymisation.yaml"
    texte = chemin.read_text(encoding="utf-8")
    assert "moteur: presidio\n" in texte
    nouveau = "" if moteur is None else f"moteur: {moteur}\n"
    chemin.write_text(texte.replace("moteur: presidio\n", nouveau), encoding="utf-8")
    rapport = construire_rapport(racine, FournisseurFactice())
    assert type(_pseudonymiseur_du_rapport(rapport)) is (attendu or PseudonymiseurPresidio)


def test_moteur_de_pseudonymisation_inconnu_refuse(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("USINE_CLE_PSEUDO", "cle-de-test-fictive")
    racine = preparer_racine(tmp_path)
    _passer_en_palier_2(racine)
    chemin = racine / "config" / "pseudonymisation.yaml"
    texte = chemin.read_text(encoding="utf-8")
    chemin.write_text(texte.replace("moteur: presidio", "moteur: spacy"), encoding="utf-8")
    with pytest.raises(UsineErreur, match="moteur"):
        construire_rapport(racine, FournisseurFactice())
    assert not (racine / "donnees" / "pseudonymes.sqlite").exists()
