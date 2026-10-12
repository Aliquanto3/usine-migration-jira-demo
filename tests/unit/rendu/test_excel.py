"""Excel du Rapport : onglet Synthèse, liens Jira, propriétés fixes (Story 0.6)."""

import io
from datetime import UTC, datetime
from pathlib import Path

import pytest
import time_machine
from openpyxl import load_workbook  # type: ignore[import-untyped]

from tests.unit.chaine.environnement import DATE_ARRETE
from tests.unit.rendu.conftest import REGLE_JAUNE, REGLE_ROUGE
from usine.config import charger_yaml
from usine.consultation.passage import ConsultationPassage
from usine.contrats.config import ConfigJira
from usine.contrats.meteo import CodeIndicateur, Couleur, Operateur, RegleDeclenchante
from usine.rendu.commun import formater_regle, tickets_en_cause
from usine.rendu.excel import EN_TETES, NOM_ONGLET, rendre_excel

NOM = f"rapport-{DATE_ARRETE.isoformat()}.xlsx"
LIBELLES = {
    Couleur.vert: "Vert",
    Couleur.jaune: "Jaune",
    Couleur.orange: "Orange",
    Couleur.rouge: "Rouge",
}


def _lignes(octets: bytes) -> list[list[tuple[object, str | None]]]:
    classeur = load_workbook(io.BytesIO(octets))
    assert classeur.sheetnames == [NOM_ONGLET]
    return [
        [(c.value, c.hyperlink.target if c.hyperlink else None) for c in ligne]
        for ligne in classeur[NOM_ONGLET].iter_rows()
    ]


def test_synthese(passage: Path) -> None:
    consultation = ConsultationPassage(passage)
    lignes = _lignes((passage / NOM).read_bytes())
    assert [valeur for valeur, _ in lignes[0][: len(EN_TETES)]] == list(EN_TETES)
    meteos = consultation.meteo()
    assert len(lignes) == len(meteos) + 1
    for ligne, meteo in zip(lignes[1:], meteos, strict=True):
        assert (ligne[0][0], ligne[1][0]) == (meteo.entite, LIBELLES[meteo.couleur])
        assert ligne[2][0] == ("\n".join(formater_regle(r) for r in meteo.regles) or None)


REGLE_INFINIE = RegleDeclenchante(
    code=CodeIndicateur.ind_3,
    couleur=Couleur.orange,
    operateur=Operateur.superieur_ou_egal,
    numerateur=2,
    denominateur=0,
    seuil_numerateur=1,
    seuil_denominateur=2,
    tickets_en_cause=(),
)


@pytest.mark.parametrize(
    ("regle", "attendu"),
    [
        (REGLE_ROUGE, "IND-1 : 7 (seuil ≥ 3)"),
        (REGLE_JAUNE, "IND-2 : 0,4 (seuil > 0,3)"),
        (REGLE_INFINIE, "IND-3 : ∞ (seuil ≥ 0,5)"),
    ],
)
def test_formater_regle(regle: RegleDeclenchante, attendu: str) -> None:
    assert formater_regle(regle) == attendu


def test_couleurs_des_cellules(passage_jaune_rouge: ConsultationPassage) -> None:
    """Teinte et texte de la cellule Météo : texte foncé sur le jaune, blanc sur le rouge."""
    onglet = load_workbook(io.BytesIO(rendre_excel(passage_jaune_rouge, DATE_ARRETE)))[NOM_ONGLET]
    attendus = {2: ("Jaune", "F9A825", "212121"), 3: ("Rouge", "C62828", "FFFFFF")}
    for ligne, (libelle, teinte, texte) in attendus.items():
        cellule = onglet.cell(row=ligne, column=2)
        assert cellule.value == libelle
        assert cellule.fill.fill_type == "solid"
        assert cellule.fill.start_color.rgb[-6:] == teinte
        assert cellule.font.color.rgb[-6:] == texte
    assert onglet.cell(row=2, column=3).value == "IND-2 : 0,4 (seuil > 0,3)"
    assert onglet.cell(row=2, column=3).alignment.wrap_text is True


def test_liens_jira(passage: Path) -> None:
    # Le Passage est publié avec `config/jira.yaml` du dépôt : ses liens suivent `url_site`.
    url_site = charger_yaml(
        Path(__file__).resolve().parents[3] / "config" / "jira.yaml", ConfigJira
    ).url_site
    consultation = ConsultationPassage(passage, url_jira=url_site)
    lignes = _lignes((passage / NOM).read_bytes())
    rouges = [m for m in consultation.meteo() if m.couleur is Couleur.rouge]
    assert rouges, "le jeu graine 42 doit avoir une Entité rouge (Scénarios de la Story 1.5)"
    vus = 0
    for ligne, meteo in zip(lignes[1:], consultation.meteo(), strict=True):
        cles = tickets_en_cause(meteo.regles)
        cellules = [c for c in ligne[len(EN_TETES) - 1 :] if c[0] is not None]
        assert [valeur for valeur, _ in cellules] == list(cles)
        for cle, lien in cellules:
            assert isinstance(cle, str)
            assert lien == consultation.url_ticket(cle)
            assert lien == f"{url_site}/browse/{cle}"
            vus += 1
    assert vus > 0


@pytest.mark.parametrize("nom_passage", ["passage_repli", "passage_sans_jira"])
def test_cles_sans_lien(nom_passage: str, request: pytest.FixtureRequest) -> None:
    """Mode repli, ou pas de `config/jira.yaml` : Passage publié, clés sans hyperlien."""
    repertoire: Path = request.getfixturevalue(nom_passage)
    lignes = _lignes((repertoire / NOM).read_bytes())
    cles = [c for ligne in lignes[1:] for c in ligne[len(EN_TETES) - 1 :] if c[0] is not None]
    assert cles
    assert all(lien is None for _, lien in cles)


def test_proprietes_fixes(passage: Path) -> None:
    proprietes = load_workbook(io.BytesIO((passage / NOM).read_bytes())).properties
    minuit = datetime(2026, 10, 9, tzinfo=UTC).replace(tzinfo=None)  # relu sans fuseau
    assert (proprietes.created, proprietes.modified) == (minuit, minuit)
    assert proprietes.creator == "Usine de migration (démonstrateur)"
    assert proprietes.title == "Rapport quotidien du 09/10/2026"


def test_rendu_deterministe(consultation: ConsultationPassage) -> None:
    """Mêmes octets, rendus à des jours d'écart."""
    with time_machine.travel(datetime(2026, 10, 12, 8, tzinfo=UTC), tick=False):
        premier = rendre_excel(consultation, DATE_ARRETE)
    with time_machine.travel(datetime(2026, 11, 3, 17, 45, tzinfo=UTC), tick=False):
        second = rendre_excel(consultation, DATE_ARRETE)
    assert premier == second


def test_sans_meteo(tmp_path: Path) -> None:
    """Passage sans Météo : en-têtes seuls, à la Date d'arrêté donnée."""
    octets = rendre_excel(ConsultationPassage(tmp_path), DATE_ARRETE)
    lignes = _lignes(octets)
    assert len(lignes) == 1
    assert [valeur for valeur, _ in lignes[0]] == list(EN_TETES)
    proprietes = load_workbook(io.BytesIO(octets)).properties
    assert proprietes.title == "Rapport quotidien du 09/10/2026"
