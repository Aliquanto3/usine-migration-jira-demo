"""PowerPoint du Rapport en 7 sections : ordre, SM-1 slide par slide, Palier 1 et 2, liens (5.4).

Le test SM-1 recompose chaque slide depuis les lectures de `ConsultationPassage` sur le même
Passage, nombres écrits par `contrats.format` : chaque texte doit être égal.
"""

import io
from dataclasses import dataclass, field, replace
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

import pytest
import time_machine
from pptx import Presentation

from tests.unit.chaine.environnement import DATE_ARRETE, PROGRAMME_REDUIT, RACINE_DEPOT
from tests.unit.rendu.conftest import (
    CONTEXTE_FICTIF,
    GRILLE,
    REGLE_JAUNE,
    REGLE_ROUGE,
    URL_JIRA,
)
from usine.config import charger_yaml
from usine.consultation.passage import ConsultationPassage
from usine.contrats.agent import AnomalieResume, ResultatListAnomalies
from usine.contrats.calcul import PERIMETRE_PROGRAMME
from usine.contrats.calendrier import jour_ouvre_precedent
from usine.contrats.config import ConfigProgramme
from usine.contrats.erreurs import UsineErreur
from usine.contrats.format import formater_date, formater_entier
from usine.contrats.llm import SortieExplication, StatutControle
from usine.contrats.meteo import (
    CodeIndicateur,
    Couleur,
    MeteoEntite,
    Operateur,
    RegleDeclenchante,
)
from usine.contrats.passage import EnveloppeFichier
from usine.contrats.prevision import ContenuPrevisions, RisqueEntreeRouge
from usine.contrats.tickets import LIBELLES, Gravite, Phase, StatutAnomalie
from usine.rendu.commun import (
    ContexteRendu,
    depuis_manifeste,
    formater_kpi,
    formater_regle,
    formater_risque,
    formater_seuils,
    libelle_kpi,
    lignes_tracabilite,
    texte_prevision,
    tickets_en_cause,
)
from usine.rendu.powerpoint import rendre_powerpoint

NOM = f"rapport-{DATE_ARRETE.isoformat()}.pptx"
GABARIT = RACINE_DEPOT / "gabarits" / "rapport.pptx"
COULEURS = {
    Couleur.vert: "Vert",
    Couleur.jaune: "Jaune",
    Couleur.orange: "Orange",
    Couleur.rouge: "Rouge",
}
MENTION = "commentaire non généré (Palier 1)"
CODE_KO = "taux_ko_cas_de_test"
TITRES_FIXES_DEBUT = ["Rapport quotidien du 09/10/2026", "Synthèse", "Carte météo du 09/10/2026"]
TITRES_FIXES_FIN = [
    "Tendances sur 8 semaines",
    "Les 10 bloquantes les plus anciennes",
    "Annexe méthode",
]


@dataclass
class SlideLue:
    """Textes d'une slide : titre, zones de texte (hors titre) et tableaux, dans l'ordre."""

    titre: str
    zones: list[str] = field(default_factory=list)
    tableaux: list[list[list[str]]] = field(default_factory=list)
    liens: dict[str, str] = field(default_factory=dict)
    """Texte de chaque run lié → adresse du lien."""


def _lire(octets: bytes) -> list[SlideLue]:
    lues: list[SlideLue] = []
    for slide in Presentation(io.BytesIO(octets)).slides:
        titre = slide.shapes.title
        lue = SlideLue(titre.text if titre is not None else "")
        for forme in slide.shapes:
            cadres = []
            if forme.has_text_frame and (titre is None or forme.shape_id != titre.shape_id):
                lue.zones.append(forme.text_frame.text)
                cadres.append(forme.text_frame)
            if forme.has_table:
                lignes = forme.table.rows
                lue.tableaux.append([[cellule.text for cellule in ligne.cells] for ligne in lignes])
                cadres.extend(cellule.text_frame for ligne in lignes for cellule in ligne.cells)
            for cadre in cadres:
                for paragraphe in cadre.paragraphs:
                    for run in paragraphe.runs:
                        if run.hyperlink.address is not None:
                            lue.liens[run.text] = run.hyperlink.address
        lues.append(lue)
    return lues


def _publie(passage: Path) -> list[SlideLue]:
    return _lire((passage / NOM).read_bytes())


def _entites(nombre: int) -> str:
    return f"{formater_entier(nombre)} Entité{'s' if nombre > 1 else ''}"


def _rouges_oranges(consultation: ConsultationPassage) -> list[MeteoEntite]:
    return [m for m in consultation.meteo() if m.couleur in (Couleur.rouge, Couleur.orange)]


# --- Gabarit, ordre et Titre ---


def test_gabarit_neutre() -> None:
    presentation = Presentation(str(GABARIT))
    assert (presentation.slide_width, presentation.slide_height) == (12_192_000, 6_858_000)
    assert len(presentation.slides) == 0
    assert presentation.core_properties.author == "Usine de migration (démonstrateur)"


def test_ordre_des_slides(passage: Path, consultation: ConsultationPassage) -> None:
    """Titre, Synthèse, Carte météo, une slide par Entité rouge ou orange, Tendances,
    bloquantes, Annexe (H-10)."""
    entites = _rouges_oranges(consultation)
    assert [m.entite for m in entites] == ["ENT-B"]  # graine 42, flux scripté : B seule rouge
    titres = [slide.titre for slide in _publie(passage)]
    assert titres == [
        *TITRES_FIXES_DEBUT,
        *(f"{m.entite} — {COULEURS[m.couleur]}" for m in entites),
        *TITRES_FIXES_FIN,
    ]


def test_titre(passage: Path, consultation: ConsultationPassage) -> None:
    programme = charger_yaml(PROGRAMME_REDUIT, ConfigProgramme)
    meteos = consultation.meteo()
    slide = _publie(passage)[0]
    assert slide.zones == [
        "\n".join(
            (
                f"Programme : {programme.nom}",
                "Date d'arrêté : 09/10/2026",
                f"Périmètre : {_entites(len(meteos))} : {', '.join(m.entite for m in meteos)}",
                "Palier 1",
            )
        )
    ]


# --- SM-1, slide par slide ---


def test_synthese(passage: Path, consultation: ConsultationPassage) -> None:
    meteos = consultation.meteo()
    slide = _publie(passage)[1]
    (tableau,) = slide.tableaux
    attendu = [["Couleur", "Nombre", "Entités"]]
    for couleur in (Couleur.rouge, Couleur.orange, Couleur.jaune, Couleur.vert):
        entites = [m.entite for m in meteos if m.couleur is couleur]
        attendu.append([COULEURS[couleur], formater_entier(len(entites)), ", ".join(entites)])
    assert tableau == attendu

    evolutions = {e.entite: e for e in consultation.evolution_meteo()}
    j1 = [
        f"{m.entite} : {COULEURS[avant]} → {COULEURS[m.couleur]}"
        for m in meteos
        if (avant := evolutions[m.entite].couleur_j1) is not None and avant is not m.couleur
    ]
    j5 = [
        f"{m.entite} : {COULEURS[avant]} → {COULEURS[m.couleur]}"
        for m in meteos
        if (avant := evolutions[m.entite].couleur_j5) is not None and avant is not m.couleur
    ]
    assert j5 == []  # graine 42, flux scripté : B rouge depuis la semaine 2, A verte
    changements, messages = slide.zones
    assert changements == "\n".join(
        (
            "Changements de Météo",
            f"Depuis J-1 : {' ; '.join(j1) or 'aucun changement'}",
            f"Depuis J-5 : {' ; '.join(j5) or 'aucun changement'}",
        )
    )

    rouges = sum(1 for m in meteos if m.couleur is Couleur.rouge)
    oranges = sum(1 for m in meteos if m.couleur is Couleur.orange)
    programme = {
        (k["code_kpi"], k["dimension"]): formater_kpi(
            k["code_kpi"], k["numerateur"], k["denominateur"]
        )
        for k in consultation.kpi(DATE_ARRETE).iter_rows(named=True)
        if k["perimetre"] == PERIMETRE_PROGRAMME
    }
    m2 = (
        f"{_entites(len(j1))} {'a' if len(j1) == 1 else 'ont'} changé de Météo depuis J-1"
        if j1
        else "Aucune Entité n'a changé de Météo depuis J-1"
    )
    assert messages == "\n".join(
        (
            "Messages clés",
            f"{_entites(rouges)} en rouge et {formater_entier(oranges)} en orange sur"
            f" {formater_entier(len(meteos))}",
            m2,
            "Anomalies bloquantes ouvertes sur le Programme :"
            f" {programme[('anomalies_ouvertes_gravite', 'bloquante')]} ; Entrées et Sorties de"
            f" la Fenêtre glissante : {programme[('entrees_fenetre', '')]} et"
            f" {programme[('sorties_fenetre', '')]}",
        )
    )
    assert messages.splitlines()[1] == "1 Entité en rouge et 0 en orange sur 2"


def test_carte_meteo(passage: Path, consultation: ConsultationPassage) -> None:
    (tableau,) = _publie(passage)[2].tableaux
    evolutions = {e.entite: e for e in consultation.evolution_meteo()}
    fleches = {-1: "↘ amélioration", 0: "→ stable", 1: "↗ aggravation"}
    attendu = [["Entité", "Météo", "Tendance vs J-5", "Règle déclenchante principale"]]
    for meteo in consultation.meteo():
        j5 = evolutions[meteo.entite].couleur_j5
        ecart = (
            None if j5 is None else (meteo.couleur.rang > j5.rang) - (meteo.couleur.rang < j5.rang)
        )
        attendu.append(
            [
                meteo.entite,
                COULEURS[meteo.couleur],
                "– non couverte" if ecart is None else fleches[ecart],
                formater_regle(meteo.regles[0]) if meteo.couleur is not Couleur.vert else "—",
            ]
        )
    assert tableau == attendu
    assert [ligne[2] for ligne in tableau[1:]] == ["→ stable", "→ stable"]  # graine 42


def _kpi_cles(consultation: ConsultationPassage, entite: str) -> list[list[str]]:
    lignes = [["KPI clé", "Valeur"]]
    for k in consultation.kpi(DATE_ARRETE, entite).iter_rows(named=True):
        code, dimension = k["code_kpi"], k["dimension"]
        retenu = (
            code in {"anomalies_ouvertes", "entrees_fenetre", "sorties_fenetre"}
            or code in {"age_median_anomalies_ouvertes", "age_max_anomalies_ouvertes"}
            or (code == "anomalies_ouvertes_gravite" and dimension in {"bloquante", "majeure"})
            or code == "taux_ko_cas_de_test"
        )
        if retenu:
            intitule = libelle_kpi(code) + (f" — {LIBELLES[dimension]}" if dimension else "")
            lignes.append([intitule, formater_kpi(code, k["numerateur"], k["denominateur"])])
    return lignes


def test_slides_entite(passage: Path, consultation: ConsultationPassage) -> None:
    slides = _publie(passage)[3:-3]
    entites = _rouges_oranges(consultation)
    assert len(slides) == len(entites)
    for slide, meteo in zip(slides, entites, strict=True):
        regles, brouillon, prevision = slide.zones
        cles = tickets_en_cause(meteo.regles)
        assert regles == "\n".join(
            (
                "Règles déclenchantes",
                *(formater_regle(r) for r in meteo.regles),
                *([f"Tickets en cause : {', '.join(cles)}"] if cles else []),
            )
        )
        assert brouillon == f"Brouillon de commentaire\n{MENTION}"
        texte = texte_prevision(consultation, meteo.entite, DATE_ARRETE, True)
        assert prevision == f"Prévision\n{texte}"
        assert slide.tableaux == [_kpi_cles(consultation, meteo.entite)]
        assert {cle: consultation.url_ticket(cle) for cle in cles} == {
            cle: slide.liens[cle] for cle in cles
        }
    # Libellés et Phases : 7 KPI non ventilés ou par gravité, puis le taux de KO par Phase.
    assert [ligne[0] for ligne in slides[0].tableaux[0][1:4]] == [
        "Anomalies ouvertes",
        "Anomalies ouvertes par gravité — Bloquante",
        "Anomalies ouvertes par gravité — Majeure",
    ]


def test_tendances(passage: Path, consultation: ConsultationPassage) -> None:
    (tableau,) = _publie(passage)[-3].tableaux
    table = consultation.tendances()
    semaines = sorted(set(table["date"].to_list()))
    assert len(semaines) == 8
    assert tableau[0] == ["Série", *(f"Semaine au {formater_date(s)}" for s in semaines)]

    def serie(perimetre: str, code: str, dimension: str) -> list[str]:
        valeurs = {
            k["date"]: formater_kpi(code, k["numerateur"], k["denominateur"])
            for k in table.iter_rows(named=True)
            if (k["perimetre"], k["code_kpi"], k["dimension"]) == (perimetre, code, dimension)
        }
        return [valeurs.get(s, "") for s in semaines]

    entites = [m.entite for m in consultation.meteo()]
    phases = ["migration_a_blanc", "recette", "tnr"]
    bloquantes = "anomalies_ouvertes_gravite", "bloquante"
    assert tableau[1:] == [
        ["Bloquantes ouvertes — Programme", *serie(PERIMETRE_PROGRAMME, *bloquantes)],
        *([f"Bloquantes ouvertes — {e}", *serie(e, *bloquantes)] for e in entites),
        *(
            [f"Taux de KO — Programme — {LIBELLES[p]}", *serie(PERIMETRE_PROGRAMME, CODE_KO, p)]
            for p in phases
        ),
    ]


def test_bloquantes_les_plus_anciennes(passage: Path, consultation: ConsultationPassage) -> None:
    slide = _publie(passage)[-2]
    resultat = consultation.anomalies(DATE_ARRETE, gravite=Gravite.bloquante, palier=1)
    anciennes = sorted(resultat.anomalies, key=lambda a: -a.age_jours_ouvres)[:10]
    assert anciennes
    assert slide.zones == [
        f"{len(anciennes)} sur {formater_entier(resultat.total)} bloquantes ouvertes"
    ]
    (tableau,) = slide.tableaux
    assert tableau == [
        ["Clé", "Entité", "Phase", "Domaine", "Statut", "Âge (Jours ouvrés)"],
        *(
            [
                a.cle_ticket,
                a.entite,
                LIBELLES[a.phase],
                a.domaine,
                LIBELLES[a.statut],
                formater_entier(a.age_jours_ouvres),
            ]
            for a in anciennes
        ),
    ]
    ages = [a.age_jours_ouvres for a in anciennes]
    assert ages == sorted(ages, reverse=True)
    for anomalie in anciennes:
        assert slide.liens[anomalie.cle_ticket] == consultation.url_ticket(anomalie.cle_ticket)
        assert slide.liens[anomalie.cle_ticket].startswith(URL_JIRA)


def test_annexe_methode(passage: Path, consultation: ConsultationPassage) -> None:
    manifeste = consultation.manifeste()
    grille, tracabilite = _publie(passage)[-1].tableaux
    assert grille == [
        ["Code", "Indicateur", "Seuils", "Version de la Grille"],
        *(
            [
                c.value,
                GRILLE.indicateurs[c].libelle,
                formater_seuils(GRILLE, c),
                manifeste.version_grille,
            ]
            for c in CodeIndicateur
            if c in GRILLE.indicateurs
        ),
    ]
    assert len(grille) == 9
    assert tracabilite == [
        ["Traçabilité", "Valeur"],
        ["Date d'arrêté", formater_date(manifeste.date_arrete)],
        ["Palier", f"Palier {manifeste.palier}"],
        ["Source", manifeste.source],
        ["Horodatage d'extraction", "sans extraction Jira"],
        ["Nombre de tickets", formater_entier(manifeste.nb_tickets)],
        ["Version de la Grille", manifeste.version_grille],
        ["Version des données", manifeste.version_donnees],
        ["Version du prompt", "sans objet (Palier 1)"],
        ["Version du modèle", "sans objet (Palier 1)"],
        ["Version des contrats", manifeste.contrats_version],
    ]
    contexte = depuis_manifeste(manifeste, GRILLE)
    assert tracabilite[1:] == [
        list(ligne) for ligne in lignes_tracabilite(consultation, DATE_ARRETE, contexte)
    ]


# --- Palier 2, liens, cas limites ---


def test_palier_2_brouillons(passage_palier_2: Path) -> None:
    consultation = ConsultationPassage(passage_palier_2, url_jira=URL_JIRA)
    contexte = depuis_manifeste(consultation.manifeste(), GRILLE)
    assert contexte.palier == 2
    slides = _lire(rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, contexte))
    brouillons = {slide.titre.split(" — ")[0]: slide.zones[1] for slide in slides[3:-3]}
    assert brouillons == {
        "ENT-B": (
            "Brouillon de commentaire\nbrouillon IA, à valider : Relancer l'équipe paramétrage."
        ),
    }
    assert slides[0].zones[0].endswith("Palier 2")


REGLE_ORANGE = RegleDeclenchante(
    code=CodeIndicateur.ind_4,
    couleur=Couleur.orange,
    operateur=Operateur.superieur,
    numerateur=1,
    denominateur=4,
    seuil_numerateur=1,
    seuil_denominateur=5,
    tickets_en_cause=(),
)


def _meteo(jour: date, entite: str, couleur: Couleur) -> MeteoEntite:
    regles = {
        Couleur.jaune: (REGLE_JAUNE,),
        Couleur.orange: (REGLE_ORANGE,),
        Couleur.rouge: (REGLE_ROUGE,),
    }.get(couleur, ())
    return MeteoEntite(
        date_arrete=jour, entite=entite, couleur=couleur, regles=regles, version_grille="abc"
    )


def _ecrire_enveloppe[T](chemin: Path, contenu: T) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    enveloppe = EnveloppeFichier[T](
        passage_id="20261012T073000Z-abcdef", date_arrete=DATE_ARRETE, contenu=contenu
    )
    chemin.write_text(enveloppe.model_dump_json(), encoding="utf-8")


def test_carte_meteo_aggravation(tmp_path: Path) -> None:
    """ENT-X jaune à J-5 et rouge à J : flèche « ↗ aggravation » (perdue en contrats 1.2)."""
    j5 = jour_ouvre_precedent(DATE_ARRETE, 5)
    _ecrire_enveloppe(
        tmp_path / "meteo.json",
        (_meteo(j5, "ENT-X", Couleur.jaune), _meteo(DATE_ARRETE, "ENT-X", Couleur.rouge)),
    )
    consultation = ConsultationPassage(tmp_path, date_arrete=DATE_ARRETE)
    slides = _lire(
        rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE_FICTIF)
    )
    (tableau,) = slides[2].tableaux
    assert tableau[1] == ["ENT-X", "Rouge", "↗ aggravation", formater_regle(REGLE_ROUGE)]


def test_palier_2_sans_brouillon(tmp_path: Path) -> None:
    """Palier 2, ENT-R rouge et ENT-O orange, Brouillon pour ENT-R seule : ENT-O « aucun »."""
    _ecrire_enveloppe(
        tmp_path / "meteo.json",
        (_meteo(DATE_ARRETE, "ENT-R", Couleur.rouge), _meteo(DATE_ARRETE, "ENT-O", Couleur.orange)),
    )
    _ecrire_enveloppe(
        tmp_path / "llm" / "explications.json",
        (
            SortieExplication(
                date_arrete=DATE_ARRETE,
                entite="ENT-R",
                texte="DEMO-6 bloque la recette de ENT-R.",
                citations=("DEMO-6",),
                brouillon_commentaire="Relancer l'équipe paramétrage.",
                statut_controle=StatutControle.conforme,
            ),
            SortieExplication(
                date_arrete=DATE_ARRETE,
                entite="ENT-O",
                texte="DEMO-24 retarde la recette de ENT-O.",
                citations=("DEMO-24",),
                brouillon_commentaire="",
                statut_controle=StatutControle.conforme,
            ),
        ),
    )
    consultation = ConsultationPassage(tmp_path, date_arrete=DATE_ARRETE)
    contexte = replace(CONTEXTE_FICTIF, palier=2)
    slides = _lire(rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, contexte))
    brouillons = {slide.titre: slide.zones[1] for slide in slides[3:-3]}
    assert brouillons == {
        "ENT-R — Rouge": (
            "Brouillon de commentaire\nbrouillon IA, à valider : Relancer l'équipe paramétrage."
        ),
        "ENT-O — Orange": "Brouillon de commentaire\naucun Brouillon de commentaire",
    }


def test_liens_et_mode_repli(passage: Path, passage_repli: Path) -> None:
    liens = {cle: url for slide in _publie(passage) for cle, url in slide.liens.items()}
    assert liens
    assert all(url == f"{URL_JIRA}/browse/{cle}" for cle, url in liens.items())
    assert all(not slide.liens for slide in _publie(passage_repli))
    assert [s.titre for s in _publie(passage_repli)] == [s.titre for s in _publie(passage)]


def test_passage_vide(tmp_path: Path) -> None:
    """Sans Météo : 6 slides, tableaux à en-têtes seuls, messages « aucune »."""
    octets = rendre_powerpoint(
        ConsultationPassage(tmp_path), GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE_FICTIF
    )
    slides = _lire(octets)
    assert [s.titre for s in slides] == [*TITRES_FIXES_DEBUT, *TITRES_FIXES_FIN]
    assert slides[0].zones == [
        "Programme : non renseigné\nDate d'arrêté : 09/10/2026\nPérimètre : aucune Entité\nPalier 1"
    ]
    assert slides[1].tableaux == [[["Couleur", "Nombre", "Entités"]]]
    assert slides[1].zones[1] == "\n".join(
        (
            "Messages clés",
            "Aucune Entité évaluée à la Date d'arrêté",
            "Aucune Entité n'a changé de Météo depuis J-1",
            "Anomalies bloquantes ouvertes sur le Programme : aucune mesure ; Entrées et Sorties"
            " de la Fenêtre glissante : aucune mesure et aucune mesure",
        )
    )
    assert slides[2].tableaux == [
        [["Entité", "Météo", "Tendance vs J-5", "Règle déclenchante principale"]]
    ]
    assert slides[3].tableaux == [[["Série"]]]
    assert slides[4].zones == ["0 sur 0 bloquantes ouvertes"]
    assert len(slides[4].tableaux[0]) == 1


def _passage_jaune(tmp_path: Path) -> ConsultationPassage:
    meteos = (
        MeteoEntite(
            date_arrete=DATE_ARRETE, entite="ENT-J", couleur=Couleur.jaune,
            regles=(REGLE_JAUNE,), version_grille="abc",
        ),
    )  # fmt: skip
    enveloppe = EnveloppeFichier[tuple[MeteoEntite, ...]](
        passage_id="20261012T073000Z-abcdef", date_arrete=DATE_ARRETE, contenu=meteos
    )
    (tmp_path / "meteo.json").write_text(enveloppe.model_dump_json(), encoding="utf-8")
    return ConsultationPassage(tmp_path, date_arrete=DATE_ARRETE)


def test_aucune_entite_rouge_ni_orange(tmp_path: Path) -> None:
    consultation = _passage_jaune(tmp_path)
    slides = _lire(
        rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE_FICTIF)
    )
    assert [s.titre for s in slides] == [*TITRES_FIXES_DEBUT, *TITRES_FIXES_FIN]
    carte = slides[2].tableaux[0]
    assert carte[1] == ["ENT-J", "Jaune", "– non couverte", formater_regle(REGLE_JAUNE)]
    assert slides[1].zones[1].splitlines()[1] == "0 Entité en rouge et 0 en orange sur 1"


REGLE_ORANGE = RegleDeclenchante(
    code=CodeIndicateur.ind_1,
    couleur=Couleur.orange,
    operateur=Operateur.superieur_ou_egal,
    numerateur=5,
    denominateur=1,
    seuil_numerateur=3,
    seuil_denominateur=1,
    tickets_en_cause=("DEMO-7",),
)


def test_slide_entite_orange_et_valeur_projetee(tmp_path: Path) -> None:
    """Orange, rouge, jaune : slides d'Entité orange puis rouge ; Risque avec valeur projetée."""
    meteos = tuple(
        MeteoEntite(
            date_arrete=DATE_ARRETE, entite=entite, couleur=couleur,
            regles=(regle,), version_grille="abc",
        )
        for entite, couleur, regle in (
            ("ENT-O", Couleur.orange, REGLE_ORANGE),
            ("ENT-R", Couleur.rouge, REGLE_ROUGE),
            ("ENT-J", Couleur.jaune, REGLE_JAUNE),
        )
    )  # fmt: skip
    risque = RisqueEntreeRouge(
        date_arrete=DATE_ARRETE,
        entite="ENT-O",
        indicateur=CodeIndicateur.ind_1,
        date_franchissement_estimee=date(2026, 10, 14),
        horizon_jours_ouvres=5,
        numerateur=5,
        denominateur=1,
        seuil_numerateur=10,
        seuil_denominateur=1,
        version_grille="abc",
        valeur_projetee_numerateur=21,
        valeur_projetee_denominateur=2,
    )
    passage_id = "20261012T073000Z-abcdef"
    (tmp_path / "meteo.json").write_text(
        EnveloppeFichier[tuple[MeteoEntite, ...]](
            passage_id=passage_id, date_arrete=DATE_ARRETE, contenu=meteos
        ).model_dump_json(),
        encoding="utf-8",
    )
    (tmp_path / "previsions.json").write_text(
        EnveloppeFichier[ContenuPrevisions](
            passage_id=passage_id,
            date_arrete=DATE_ARRETE,
            contenu=ContenuPrevisions(previsions=(), risques=(risque,)),
        ).model_dump_json(),
        encoding="utf-8",
    )
    consultation = ConsultationPassage(tmp_path, date_arrete=DATE_ARRETE)
    slides = _lire(
        rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE_FICTIF)
    )
    assert [s.titre for s in slides] == [
        *TITRES_FIXES_DEBUT,
        "ENT-O — Orange",
        "ENT-R — Rouge",
        *TITRES_FIXES_FIN,
    ]
    texte = formater_risque(risque)
    assert "valeur projetée 10,5" in texte
    assert slides[3].zones[2] == f"Prévision\n{texte}"


class _ConsultationBloquantes(ConsultationPassage):
    """Passage vide dont `anomalies` rend des bloquantes fixées, dans l'ordre des clés."""

    AGES = (3, 15, 7, 20, 7, 1, 30, 12, 9, 25, 4, 18)

    def anomalies(self, *args: Any, **kwargs: Any) -> ResultatListAnomalies:
        anomalies = tuple(
            AnomalieResume(
                cle_ticket=f"DEMO-{rang:02d}",
                entite="ENT-A",
                gravite=Gravite.bloquante,
                phase=Phase.recette,
                domaine="Paie",
                statut=StatutAnomalie.nouvelle,
                age_jours_ouvres=age,
            )
            for rang, age in enumerate(self.AGES, start=1)
        )
        return ResultatListAnomalies(
            date_arrete=DATE_ARRETE, anomalies=anomalies, total=len(anomalies)
        )


def test_bloquantes_limite_a_10(tmp_path: Path) -> None:
    """12 bloquantes : les 10 plus anciennes, âge décroissant, ex aequo dans l'ordre lu."""
    consultation = _ConsultationBloquantes(tmp_path, url_jira=URL_JIRA)
    slide = _lire(
        rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE_FICTIF)
    )[-2]
    assert slide.zones == ["10 sur 12 bloquantes ouvertes"]
    (tableau,) = slide.tableaux
    assert [(ligne[0], ligne[5]) for ligne in tableau[1:]] == [
        ("DEMO-07", "30"),
        ("DEMO-10", "25"),
        ("DEMO-04", "20"),
        ("DEMO-12", "18"),
        ("DEMO-02", "15"),
        ("DEMO-08", "12"),
        ("DEMO-09", "9"),
        ("DEMO-03", "7"),
        ("DEMO-05", "7"),
        ("DEMO-11", "4"),
    ]
    assert slide.liens == {ligne[0]: f"{URL_JIRA}/browse/{ligne[0]}" for ligne in tableau[1:]}


def test_couleurs_des_cellules(passage_jaune_rouge: ConsultationPassage) -> None:
    """Teinte et texte de la cellule Météo : texte foncé sur le jaune, blanc sur le rouge."""
    octets = rendre_powerpoint(
        passage_jaune_rouge, GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE_FICTIF
    )
    slide = Presentation(io.BytesIO(octets)).slides[2]
    tableau = next(forme.table for forme in slide.shapes if forme.has_table)
    attendus = {1: ("Jaune", "F9A825", "212121"), 2: ("Rouge", "C62828", "FFFFFF")}
    for ligne, (libelle, teinte, texte) in attendus.items():
        cellule = tableau.cell(ligne, 1)
        assert cellule.text == libelle
        assert str(cellule.fill.fore_color.rgb) == teinte
        assert str(cellule.text_frame.paragraphs[0].runs[0].font.color.rgb) == texte


# --- Documents ---


def test_proprietes_fixes(passage: Path) -> None:
    proprietes = Presentation(str(passage / NOM)).core_properties
    minuit = datetime(2026, 10, 9, tzinfo=UTC).replace(tzinfo=None)  # relu sans fuseau
    assert (proprietes.created, proprietes.modified) == (minuit, minuit)
    assert proprietes.author == "Usine de migration (démonstrateur)"
    assert proprietes.title == "Rapport quotidien du 09/10/2026"


@pytest.mark.parametrize("palier", [1, 2])
def test_rendu_deterministe(passage_palier_2: Path, palier: int) -> None:
    consultation = ConsultationPassage(passage_palier_2, url_jira=URL_JIRA)
    contexte = ContexteRendu(
        palier=2 if palier == 2 else 1, source="journal", version_grille="g",
        version_donnees="d", grille=GRILLE, programme="Programme fictif",
    )  # fmt: skip
    gabarit = GABARIT.read_bytes()
    with time_machine.travel(datetime(2026, 10, 12, 8, tzinfo=UTC), tick=False):
        premier = rendre_powerpoint(consultation, gabarit, DATE_ARRETE, contexte)
    with time_machine.travel(datetime(2026, 11, 3, 17, 45, tzinfo=UTC), tick=False):
        second = rendre_powerpoint(consultation, gabarit, DATE_ARRETE, contexte)
    assert premier == second


def test_gabarit_illisible(passage: Path) -> None:
    with pytest.raises(UsineErreur, match="gabarit"):
        rendre_powerpoint(
            ConsultationPassage(passage), b"pas un pptx", DATE_ARRETE, CONTEXTE_FICTIF
        )
