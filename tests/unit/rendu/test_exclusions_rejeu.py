"""Garde des repères de rendu de `EXCLUSIONS_REJEU` (contrats 1.2, AD-7, décision C8).

`contrats` n'importe pas `rendu` : ce test relie `rapport_xlsx` (adresse de cellule) et
`rapport_pptx` (libellé de ligne) aux lignes « Horodatage d'extraction » et « Version des
données » de `LIBELLES_TRACABILITE`, et vérifie que les rendus produits y affichent bien les
champs exclus du Rejeu au manifeste.
"""

import io
from datetime import UTC, datetime

from openpyxl import load_workbook  # type: ignore[import-untyped]
from pptx import Presentation

from tests.unit.chaine.environnement import DATE_ARRETE, RACINE_DEPOT
from usine.consultation.passage import ConsultationPassage
from usine.contrats.passage import EXCLUSIONS_REJEU
from usine.rendu.commun import LIBELLES_TRACABILITE, ContexteRendu
from usine.rendu.excel import rendre_excel
from usine.rendu.powerpoint import rendre_powerpoint

GABARIT = RACINE_DEPOT / "gabarits" / "rapport.pptx"
EXCLUS = ("Horodatage d'extraction", "Version des données")
"""Libellés des lignes de traçabilité qui affichent `horodatage_extraction` et
`version_donnees`, exclus du Rejeu au manifeste."""
INSTANT = datetime(2026, 10, 12, 7, 5, 9, tzinfo=UTC)
CONTEXTE = ContexteRendu(
    palier=1,
    source="composite",
    version_grille="grille-1",
    version_donnees="donnees-1",
    horodatage_extraction=INSTANT,
)
AFFICHES = {
    "Horodatage d'extraction": "12/10/2026 09:05:09",
    "Version des données": "donnees-1",
}


def _repere(cle: str) -> tuple[str, str]:
    onglet, _, repere = cle.partition("!")
    return onglet, repere


def test_libelles_exclus_dans_la_tracabilite() -> None:
    assert set(EXCLUS) <= set(LIBELLES_TRACABILITE)


def test_rapport_xlsx_designe_les_cellules_exclues(consultation: ConsultationPassage) -> None:
    """`Traçabilité!B5` et `!B8` : valeur (colonne B) des lignes exclues, en-têtes en ligne 1."""
    attendus = {
        ("Traçabilité", f"B{LIBELLES_TRACABILITE.index(libelle) + 2}") for libelle in EXCLUS
    }
    assert {_repere(cle) for cle in EXCLUSIONS_REJEU["rapport_xlsx"]} == attendus
    classeur = load_workbook(io.BytesIO(rendre_excel(consultation, DATE_ARRETE, CONTEXTE)))
    for cle in EXCLUSIONS_REJEU["rapport_xlsx"]:
        onglet, cellule = _repere(cle)
        ligne = classeur[onglet][cellule].row
        libelle = classeur[onglet][f"A{ligne}"].value
        assert libelle in EXCLUS, cle
        assert classeur[onglet][cellule].value == AFFICHES[libelle], cle


def test_rapport_pptx_designe_les_lignes_exclues(consultation: ConsultationPassage) -> None:
    """`Annexe méthode!<libellé>` : ligne du tableau de traçabilité de la dernière slide."""
    assert {_repere(cle) for cle in EXCLUSIONS_REJEU["rapport_pptx"]} == {
        ("Annexe méthode", libelle) for libelle in EXCLUS
    }
    octets = rendre_powerpoint(consultation, GABARIT.read_bytes(), DATE_ARRETE, CONTEXTE)
    slides = {
        slide.shapes.title.text: slide
        for slide in Presentation(io.BytesIO(octets)).slides
        if slide.shapes.title is not None
    }
    for cle in EXCLUSIONS_REJEU["rapport_pptx"]:
        titre, libelle = _repere(cle)
        lignes = {
            ligne.cells[0].text: ligne.cells[1].text
            for forme in slides[titre].shapes
            if forme.has_table
            for ligne in forme.table.rows
        }
        assert lignes[libelle] == AFFICHES[libelle], cle
