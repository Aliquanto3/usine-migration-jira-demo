"""Excel du Rapport quotidien : onglet unique « Synthèse » (Story 0.6, amorce de FR-26).

Une ligne par Entité, dans l'ordre de `consultation.meteo()` : Entité, Météo, Règles
déclenchantes, puis une clé de ticket en cause par colonne, chacune liée à
`consultation.url_ticket(cle)` quand l'URL existe (FR-29).
"""

import io
from datetime import date

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]

from usine.contrats.meteo import LIBELLES_COULEURS
from usine.ports.consultation import Consultation
from usine.rendu.commun import (
    AUTEUR,
    TEINTES_COULEURS,
    TEXTES_COULEURS,
    archive_deterministe,
    formater_regle,
    horodatage_fixe,
    tickets_en_cause,
    titre_document,
)

NOM_ONGLET = "Synthèse"
EN_TETES = ("Entité", "Météo", "Règles déclenchantes", "Tickets en cause")
COLONNE_TICKETS = len(EN_TETES)
"""Colonne (1-indexée) de la première clé de ticket en cause ; les suivantes à sa droite."""


def rendre_excel(consultation: Consultation, date_arrete: date) -> bytes:
    """Octets du classeur Excel du Passage à la Date d'arrêté, déterministes.

    Sans Météo à cette date, le classeur ne porte que les en-têtes.
    """
    classeur = Workbook()
    onglet = classeur.active
    if onglet is None:  # pragma: no cover - un classeur neuf a toujours un onglet actif
        onglet = classeur.create_sheet()
    onglet.title = NOM_ONGLET
    gras = Font(bold=True)
    for colonne, en_tete in enumerate(EN_TETES, start=1):
        onglet.cell(row=1, column=colonne, value=en_tete).font = gras
    for ligne, meteo in enumerate(consultation.meteo(date_arrete), start=2):
        onglet.cell(row=ligne, column=1, value=meteo.entite)
        cellule = onglet.cell(row=ligne, column=2, value=LIBELLES_COULEURS[meteo.couleur])
        teinte = TEINTES_COULEURS[meteo.couleur]
        cellule.fill = PatternFill(fill_type="solid", start_color=teinte, end_color=teinte)
        cellule.font = Font(bold=True, color=TEXTES_COULEURS[meteo.couleur])
        regles = "\n".join(formater_regle(regle) for regle in meteo.regles)
        cellule = onglet.cell(row=ligne, column=3, value=regles or None)
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        for decalage, cle in enumerate(tickets_en_cause(meteo.regles)):
            cellule = onglet.cell(row=ligne, column=COLONNE_TICKETS + decalage, value=cle)
            url = consultation.url_ticket(cle)
            if url is not None:
                cellule.hyperlink = url
                cellule.style = "Hyperlink"
    onglet.column_dimensions["A"].width = 14
    onglet.column_dimensions["B"].width = 10
    onglet.column_dimensions["C"].width = 32
    onglet.freeze_panes = "A2"

    proprietes = classeur.properties
    proprietes.creator = AUTEUR
    proprietes.lastModifiedBy = AUTEUR
    proprietes.title = titre_document(date_arrete)
    proprietes.created = horodatage_fixe(date_arrete)
    proprietes.modified = horodatage_fixe(date_arrete)
    tampon = io.BytesIO()
    classeur.save(tampon)
    return archive_deterministe(tampon.getvalue(), date_arrete)
