"""PowerPoint du Rapport quotidien : 7 sections dans l'ordre de H-10 (FR-27, FR-28, FR-29, SM-1).

Construit sur le gabarit neutre `gabarits/rapport.pptx` (dispositions « Title Slide » et
« Title Only »). Chaque nombre est une valeur lue via `Consultation` (ou le `ContexteRendu`),
formatée par `contrats.format` : le rendu ne calcule rien. Compter les Entités affichées,
trier pour sélectionner et comparer deux couleurs lues sont de la mise en forme. Toute clé de
ticket affichée est liée à `consultation.url_ticket(cle)` quand l'URL existe (FR-29).

1. Titre : Programme, Date d'arrêté, périmètre, Palier.
2. Synthèse : Entités par couleur, changements depuis J-1 et J-5, trois messages clés construits
   par gabarit de phrase (jamais par LLM).
3. Carte météo : Entité, Météo, tendance vs J-5, Règle déclenchante principale.
4. Une slide par Entité rouge ou orange, dans l'ordre de `meteo(D)` : Règles, KPI clés, zone
   Brouillon de commentaire (« commentaire non généré (Palier 1) » en Palier 1), Prévision.
5. Tendances : bloquantes ouvertes (Programme puis Entités) et taux de KO du Programme par Phase,
   une colonne par semaine de `tendances()`.
6. Les 10 bloquantes les plus anciennes : par âge décroissant, puis dans l'ordre lu.
7. Annexe méthode : Grille appliquée et traçabilité (`lignes_tracabilite`).

Rejeu : l'Annexe affiche l'horodatage d'extraction et la version des données, exclus du Rejeu
au manifeste ; à déclarer sous `EXCLUSIONS_REJEU["rapport_pptx"]` (`Annexe méthode!…`, 5.5).
"""

import io
from collections.abc import Sequence
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.presentation import Presentation as Presentation_
from pptx.slide import Slide
from pptx.table import Table, _Cell
from pptx.text.text import TextFrame
from pptx.util import Emu, Length, Pt

from usine.contrats.calcul import PERIMETRE_PROGRAMME
from usine.contrats.erreurs import UsineErreur
from usine.contrats.format import formater_date, formater_entier
from usine.contrats.meteo import CodeIndicateur, Couleur, EvolutionMeteo, MeteoEntite
from usine.contrats.tickets import Gravite
from usine.ports.consultation import Consultation
from usine.rendu.commun import (
    AUTEUR,
    CODE_TAUX_KO,
    MENTION_BROUILLON,
    MENTION_PALIER_1,
    TEINTES_COULEURS,
    TEXTES_COULEURS,
    ContexteRendu,
    archive_deterministe,
    formater_kpi,
    formater_regle,
    formater_seuils,
    horodatage_fixe,
    libelle,
    libelle_couleur,
    libelle_indicateur,
    libelle_kpi,
    libelle_perimetre,
    lignes_tracabilite,
    texte_prevision,
    tickets_en_cause,
    titre_document,
)

DISPOSITION_TITRE = "Title Slide"
DISPOSITION = "Title Only"
"""Dispositions du gabarit : slide de titre, puis titre seul pour les autres slides."""

MARGE = 457_200
"""Marge latérale (1,27 cm), en EMU."""

HAUT_CONTENU = 1_500_000
"""Position verticale du contenu, sous le titre, en EMU."""

HAUTEUR_LIGNE = 330_000
HAUTEUR_LIGNE_DENSE = 250_000
TAILLE_TEXTE = Pt(14)
TAILLE_ENTITE = Pt(12)
TAILLE_TABLEAU = Pt(11)
TAILLE_DENSE = Pt(9)

TITRE_SYNTHESE = "Synthèse"
TITRE_CARTE = "Carte météo"
TITRE_TENDANCES = "Tendances sur 8 semaines"
TITRE_BLOQUANTES = "Les 10 bloquantes les plus anciennes"
TITRE_ANNEXE = "Annexe méthode"

NOMBRE_BLOQUANTES = 10
"""Bloquantes affichées sur la slide des plus anciennes."""

COULEURS_SLIDE_ENTITE = (Couleur.rouge, Couleur.orange)
"""Couleurs qui reçoivent une slide d'Entité."""

EN_TETES_SYNTHESE = ("Couleur", "Nombre", "Entités")
EN_TETES_CARTE = ("Entité", "Météo", "Tendance vs J-5", "Règle déclenchante principale")
EN_TETES_KPI = ("KPI clé", "Valeur")
EN_TETE_TENDANCES = "Série"
EN_TETES_BLOQUANTES = ("Clé", "Entité", "Phase", "Domaine", "Statut", "Âge (Jours ouvrés)")
EN_TETES_GRILLE = ("Code", "Indicateur", "Seuils", "Version de la Grille")
EN_TETES_TRACABILITE = ("Traçabilité", "Valeur")

AGGRAVATION = "↗ aggravation"
AMELIORATION = "↘ amélioration"
STABLE = "→ stable"
TENDANCE_NON_COUVERTE = "– non couverte"
SANS_REGLE = "—"
AUCUN_CHANGEMENT = "aucun changement"
AUCUNE_MESURE = "aucune mesure"
SANS_BROUILLON = "aucun Brouillon de commentaire"
PROGRAMME_INCONNU = "non renseigné"

TITRE_CHANGEMENTS = "Changements de Météo"
TITRE_MESSAGES = "Messages clés"
TITRE_REGLES = "Règles déclenchantes"
TITRE_TICKETS = "Tickets en cause : "
TITRE_BROUILLON = "Brouillon de commentaire"
TITRE_PREVISION = "Prévision"

CODE_BLOQUANTES = "anomalies_ouvertes_gravite"
"""KPI des Anomalies ouvertes par gravité, lu à la dimension `bloquante`."""

KPI_CLES: tuple[tuple[str, frozenset[str]], ...] = (
    ("anomalies_ouvertes", frozenset({""})),
    (CODE_BLOQUANTES, frozenset({Gravite.bloquante.value, Gravite.majeure.value})),
    ("entrees_fenetre", frozenset({""})),
    ("sorties_fenetre", frozenset({""})),
    ("age_median_anomalies_ouvertes", frozenset({""})),
    ("age_max_anomalies_ouvertes", frozenset({""})),
    (CODE_TAUX_KO, frozenset()),
)
"""KPI clés d'une slide d'Entité : code et dimensions retenues (vide : toutes les Phases).

Affichés dans l'ordre lu de `kpi(D, entite)`, celui du domaine.
"""


# --- Gabarit et formes ---


def _rvb(hexadecimal: str) -> RGBColor:
    rvb: RGBColor = RGBColor.from_string(hexadecimal)  # type: ignore[no-untyped-call]
    return rvb


def _gabarit(octets: bytes) -> Presentation_:
    try:
        return Presentation(io.BytesIO(octets))
    except Exception as exc:
        raise UsineErreur(f"gabarit PowerPoint illisible ({type(exc).__name__} : {exc})") from exc


def _slide(presentation: Presentation_, disposition: str, titre: str) -> Slide:
    modele = presentation.slide_layouts.get_by_name(disposition)
    if modele is None:
        raise UsineErreur(f"gabarit PowerPoint : disposition « {disposition} » absente")
    slide: Slide = presentation.slides.add_slide(modele)
    forme = slide.shapes.title
    if forme is None:
        raise UsineErreur(f"gabarit PowerPoint : disposition « {disposition} » sans titre")
    forme.text = titre
    return slide


def _largeur(presentation: Presentation_) -> int:
    return int(presentation.slide_width) - 2 * MARGE


def _police(cadre: TextFrame, taille: Length, gras: bool = False) -> None:
    for paragraphe in cadre.paragraphs:
        for run in paragraphe.runs:
            run.font.size = taille
            if gras:
                run.font.bold = True


def _zone(
    slide: Slide,
    position: tuple[int, int, int, int],
    paragraphes: Sequence[str],
    titre: str | None = None,
    taille: Length = TAILLE_TEXTE,
) -> TextFrame:
    """Zone de texte, un paragraphe par ligne, précédée d'un titre en gras s'il est donné."""
    gauche, haut, largeur, hauteur = position
    forme = slide.shapes.add_textbox(Emu(gauche), Emu(haut), Emu(largeur), Emu(hauteur))
    cadre: TextFrame = forme.text_frame
    cadre.word_wrap = True
    lignes = [*([] if titre is None else [titre]), *paragraphes]
    for rang, texte in enumerate(lignes):
        paragraphe = cadre.paragraphs[0] if rang == 0 else cadre.add_paragraph()  # type: ignore[no-untyped-call]
        run = paragraphe.add_run()
        run.text = texte
        run.font.size = taille
        if titre is not None and rang == 0:
            run.font.bold = True
    return cadre


def _tableau(
    slide: Slide,
    position: tuple[int, int, int],
    en_tetes: Sequence[str],
    lignes: Sequence[Sequence[str]],
    proportions: Sequence[int],
    dense: bool = False,
) -> Table:
    """Tableau à en-têtes en gras ; largeurs de colonnes proportionnelles à `proportions`."""
    gauche, haut, largeur = position
    hauteur_ligne = HAUTEUR_LIGNE_DENSE if dense else HAUTEUR_LIGNE
    taille = TAILLE_DENSE if dense else TAILLE_TABLEAU
    forme = slide.shapes.add_table(
        len(lignes) + 1,
        len(en_tetes),
        Emu(gauche),
        Emu(haut),
        Emu(largeur),
        Emu(hauteur_ligne * (len(lignes) + 1)),
    )
    tableau = forme.table
    total = sum(proportions)
    reste = largeur
    for rang, proportion in enumerate(proportions):
        colonne = reste if rang == len(proportions) - 1 else largeur * proportion // total
        tableau.columns[rang].width = Emu(colonne)
        reste -= colonne
    for rang, texte in enumerate(en_tetes):
        cellule = tableau.cell(0, rang)
        cellule.text = texte
        _police(cellule.text_frame, taille, gras=True)
    for ligne, valeurs in enumerate(lignes, start=1):
        for rang, texte in enumerate(valeurs):
            cellule = tableau.cell(ligne, rang)
            cellule.text = texte
            _police(cellule.text_frame, taille)
    return tableau


def _teinter(cellule: _Cell, couleur: Couleur) -> None:
    """Teinte de la couleur de Météo, texte lisible en gras."""
    cellule.fill.solid()  # type: ignore[no-untyped-call]
    cellule.fill.fore_color.rgb = _rvb(TEINTES_COULEURS[couleur])
    for run in cellule.text_frame.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = _rvb(TEXTES_COULEURS[couleur])


def _lier(cellule: _Cell, cle: str, consultation: Consultation) -> None:
    """Clé de ticket liée à `url_ticket(cle)` quand l'URL existe (FR-29)."""
    url = consultation.url_ticket(cle)
    if url is not None:
        for run in cellule.text_frame.paragraphs[0].runs:
            run.hyperlink.address = url


# --- Textes ---


def nombre_entites(nombre: int) -> str:
    """`1 Entité`, `2 Entités` (`0 Entité`)."""
    return f"{formater_entier(nombre)} Entité{'s' if nombre > 1 else ''}"


def perimetre(meteos: Sequence[MeteoEntite]) -> str:
    """`2 Entités : ENT-A, ENT-B`, ou « aucune Entité »."""
    if not meteos:
        return "aucune Entité"
    return f"{nombre_entites(len(meteos))} : {', '.join(m.entite for m in meteos)}"


def tendance(couleur: Couleur, couleur_j5: Couleur | None) -> str:
    """Flèche de la Météo J par rapport à J-5, selon l'ordre de gravité de `Couleur`."""
    if couleur_j5 is None:
        return TENDANCE_NON_COUVERTE
    if couleur > couleur_j5:
        return AGGRAVATION
    if couleur < couleur_j5:
        return AMELIORATION
    return STABLE


def changements(
    meteos: Sequence[MeteoEntite], evolutions: dict[str, EvolutionMeteo], rang: int
) -> tuple[str, ...]:
    """`ENT-A : Orange → Rouge` par Entité couverte dont la Météo diffère de J-`rang` (1 ou 5)."""
    textes: list[str] = []
    for meteo in meteos:
        evolution = evolutions.get(meteo.entite)
        if evolution is None:
            continue
        avant = evolution.couleur_j1 if rang == 1 else evolution.couleur_j5
        if avant is not None and avant is not meteo.couleur:
            textes.append(
                f"{meteo.entite} : {libelle_couleur(avant)} → {libelle_couleur(meteo.couleur)}"
            )
    return tuple(textes)


def message_couleurs(meteos: Sequence[MeteoEntite]) -> str:
    """M1 : `N Entités en rouge et M en orange sur T`."""
    if not meteos:
        return "Aucune Entité évaluée à la Date d'arrêté"
    rouges = sum(1 for m in meteos if m.couleur is Couleur.rouge)
    oranges = sum(1 for m in meteos if m.couleur is Couleur.orange)
    return (
        f"{nombre_entites(rouges)} en rouge et {formater_entier(oranges)} en orange"
        f" sur {formater_entier(len(meteos))}"
    )


def message_changements(nombre: int) -> str:
    """M2 : `K Entités ont changé de Météo depuis J-1`."""
    if nombre == 0:
        return "Aucune Entité n'a changé de Météo depuis J-1"
    verbe = "a" if nombre == 1 else "ont"
    return f"{nombre_entites(nombre)} {verbe} changé de Météo depuis J-1"


def message_programme(bloquantes: str, entrees: str, sorties: str) -> str:
    """M3 : bloquantes ouvertes, Entrées et Sorties de la Fenêtre glissante du Programme."""
    return (
        f"Anomalies bloquantes ouvertes sur le Programme : {bloquantes} ; Entrées et Sorties"
        f" de la Fenêtre glissante : {entrees} et {sorties}"
    )


def kpi_cles(consultation: Consultation, jour: date, entite: str) -> list[list[str]]:
    """Libellé et valeur des KPI clés de l'Entité, dans l'ordre lu."""
    retenus = dict(KPI_CLES)
    lignes: list[list[str]] = []
    for mesure in consultation.kpi(jour, entite).iter_rows(named=True):
        code, dimension = mesure["code_kpi"], mesure["dimension"]
        if code not in retenus or (retenus[code] and dimension not in retenus[code]):
            continue
        intitule = libelle_kpi(code) + (f" — {libelle(dimension)}" if dimension else "")
        lignes.append([intitule, formater_kpi(code, mesure["numerateur"], mesure["denominateur"])])
    return lignes


def tendances_lignes(consultation: Consultation) -> tuple[list[date], list[list[str]]]:
    """Semaines, puis lignes : bloquantes ouvertes (Programme, Entités), taux de KO Programme."""
    table = consultation.tendances()
    semaines: list[date] = sorted(set(table.get_column("date").to_list()))
    rangs = {semaine: rang for rang, semaine in enumerate(semaines, start=1)}
    bloquantes: dict[str, list[str]] = {}
    taux: dict[str, list[str]] = {}
    for mesure in table.iter_rows(named=True):
        code, dimension, perimetre_mesure = (
            mesure["code_kpi"],
            mesure["dimension"],
            mesure["perimetre"],
        )
        if code == CODE_BLOQUANTES and dimension == Gravite.bloquante.value:
            intitule = f"Bloquantes ouvertes — {libelle_perimetre(perimetre_mesure)}"
            cible = bloquantes.setdefault(perimetre_mesure, [intitule] + [""] * len(semaines))
        elif code == CODE_TAUX_KO and perimetre_mesure == PERIMETRE_PROGRAMME:
            intitule = f"Taux de KO — Programme — {libelle(dimension)}"
            cible = taux.setdefault(dimension, [intitule] + [""] * len(semaines))
        else:
            continue
        cible[rangs[mesure["date"]]] = formater_kpi(
            code, mesure["numerateur"], mesure["denominateur"]
        )
    programme = [bloquantes.pop(PERIMETRE_PROGRAMME)] if PERIMETRE_PROGRAMME in bloquantes else []
    return semaines, [*programme, *bloquantes.values(), *taux.values()]


def brouillon(consultation: Consultation, entite: str, contexte: ContexteRendu) -> str:
    """Zone Brouillon d'une slide d'Entité : mention Palier 1, Brouillon marqué, ou aucun."""
    if contexte.palier == 1:
        return MENTION_PALIER_1
    for explication in consultation.explications(entite, palier=contexte.palier):
        if explication.brouillon_commentaire:
            return f"{MENTION_BROUILLON} : {explication.brouillon_commentaire}"
    return SANS_BROUILLON


# --- Slides ---


def _titre(
    presentation: Presentation_, meteos: Sequence[MeteoEntite], jour: date, contexte: ContexteRendu
) -> None:
    slide = _slide(presentation, DISPOSITION_TITRE, titre_document(jour))
    slide.placeholders[1].text_frame.text = "\n".join(
        (
            f"Programme : {contexte.programme or PROGRAMME_INCONNU}",
            f"Date d'arrêté : {formater_date(jour)}",
            f"Périmètre : {perimetre(meteos)}",
            f"Palier {contexte.palier}",
        )
    )


def _synthese(
    presentation: Presentation_,
    consultation: Consultation,
    meteos: Sequence[MeteoEntite],
    jour: date,
) -> None:
    slide = _slide(presentation, DISPOSITION, TITRE_SYNTHESE)
    largeur = _largeur(presentation)
    couleurs = sorted(Couleur, reverse=True) if meteos else []
    lignes = []
    for couleur in couleurs:
        entites = [m.entite for m in meteos if m.couleur is couleur]
        lignes.append([libelle_couleur(couleur), formater_entier(len(entites)), ", ".join(entites)])
    tableau = _tableau(slide, (MARGE, HAUT_CONTENU, largeur), EN_TETES_SYNTHESE, lignes, (2, 1, 7))
    for ligne, couleur in enumerate(couleurs, start=1):
        _teinter(tableau.cell(ligne, 0), couleur)

    evolutions = {e.entite: e for e in consultation.evolution_meteo(jour)}
    depuis_j1 = changements(meteos, evolutions, 1)
    depuis_j5 = changements(meteos, evolutions, 5)
    _zone(
        slide,
        (MARGE, 3_250_000, largeur, 900_000),
        (
            f"Depuis J-1 : {' ; '.join(depuis_j1) or AUCUN_CHANGEMENT}",
            f"Depuis J-5 : {' ; '.join(depuis_j5) or AUCUN_CHANGEMENT}",
        ),
        titre=TITRE_CHANGEMENTS,
    )

    programme = {
        (mesure["code_kpi"], mesure["dimension"]): formater_kpi(
            mesure["code_kpi"], mesure["numerateur"], mesure["denominateur"]
        )
        for mesure in consultation.kpi(jour).iter_rows(named=True)
        if mesure["perimetre"] == PERIMETRE_PROGRAMME
    }
    _zone(
        slide,
        (MARGE, 4_300_000, largeur, 2_200_000),
        (
            message_couleurs(meteos),
            message_changements(len(depuis_j1)),
            message_programme(
                programme.get((CODE_BLOQUANTES, Gravite.bloquante.value), AUCUNE_MESURE),
                programme.get(("entrees_fenetre", ""), AUCUNE_MESURE),
                programme.get(("sorties_fenetre", ""), AUCUNE_MESURE),
            ),
        ),
        titre=TITRE_MESSAGES,
    )


def _carte(
    presentation: Presentation_,
    consultation: Consultation,
    meteos: Sequence[MeteoEntite],
    jour: date,
) -> None:
    slide = _slide(presentation, DISPOSITION, f"{TITRE_CARTE} du {formater_date(jour)}")
    evolutions = {e.entite: e for e in consultation.evolution_meteo(jour)}
    lignes: list[list[str]] = []
    for meteo in meteos:
        evolution = evolutions.get(meteo.entite)
        principale = (
            formater_regle(meteo.regles[0])
            if meteo.couleur is not Couleur.vert and meteo.regles
            else SANS_REGLE
        )
        lignes.append(
            [
                meteo.entite,
                libelle_couleur(meteo.couleur),
                tendance(meteo.couleur, None if evolution is None else evolution.couleur_j5),
                principale,
            ]
        )
    tableau = _tableau(
        slide, (MARGE, HAUT_CONTENU, _largeur(presentation)), EN_TETES_CARTE, lignes, (2, 2, 3, 7)
    )
    for ligne, meteo in enumerate(meteos, start=1):
        _teinter(tableau.cell(ligne, 1), meteo.couleur)


def _entite(
    presentation: Presentation_,
    consultation: Consultation,
    meteo: MeteoEntite,
    jour: date,
    contexte: ContexteRendu,
) -> None:
    slide = _slide(presentation, DISPOSITION, f"{meteo.entite} — {libelle_couleur(meteo.couleur)}")
    largeur = _largeur(presentation)
    colonne = largeur // 2 - 100_000
    droite = MARGE + largeur - colonne

    regles = _zone(
        slide,
        (MARGE, HAUT_CONTENU, colonne, 1_400_000),
        [formater_regle(regle) for regle in meteo.regles] or [SANS_REGLE],
        titre=TITRE_REGLES,
        taille=TAILLE_ENTITE,
    )
    cles = tickets_en_cause(meteo.regles)
    if cles:
        paragraphe = regles.add_paragraph()  # type: ignore[no-untyped-call]
        morceaux: list[tuple[str, str | None]] = [(TITRE_TICKETS, None)]
        for rang, cle in enumerate(cles):
            morceaux.extend([(", ", None)] if rang else [])
            morceaux.append((cle, consultation.url_ticket(cle)))
        for texte, url in morceaux:
            run = paragraphe.add_run()
            run.text = texte
            run.font.size = TAILLE_ENTITE
            if url is not None:
                run.hyperlink.address = url

    _tableau(
        slide,
        (MARGE, 3_000_000, colonne),
        EN_TETES_KPI,
        kpi_cles(consultation, jour, meteo.entite),
        (4, 1),
        dense=True,
    )
    _zone(
        slide,
        (droite, HAUT_CONTENU, colonne, 2_200_000),
        (brouillon(consultation, meteo.entite, contexte),),
        titre=TITRE_BROUILLON,
        taille=TAILLE_ENTITE,
    )
    prevision = texte_prevision(consultation, meteo.entite, jour, contexte.previsions_calculees)
    _zone(
        slide,
        (droite, 3_900_000, colonne, 2_600_000),
        prevision.split("\n"),
        titre=TITRE_PREVISION,
        taille=TAILLE_ENTITE,
    )


def _tendances(presentation: Presentation_, consultation: Consultation) -> None:
    slide = _slide(presentation, DISPOSITION, TITRE_TENDANCES)
    semaines, lignes = tendances_lignes(consultation)
    _tableau(
        slide,
        (MARGE, HAUT_CONTENU, _largeur(presentation)),
        (EN_TETE_TENDANCES, *(f"Semaine au {formater_date(s)}" for s in semaines)),
        lignes,
        (4, *(1 for _ in semaines)),
        dense=True,
    )


def _bloquantes(
    presentation: Presentation_, consultation: Consultation, jour: date, contexte: ContexteRendu
) -> None:
    slide = _slide(presentation, DISPOSITION, TITRE_BLOQUANTES)
    resultat = consultation.anomalies(jour, gravite=Gravite.bloquante, palier=contexte.palier)
    # Tri stable : à âge égal, l'ordre lu (par clé) est conservé.
    triees = sorted(resultat.anomalies, key=lambda a: a.age_jours_ouvres, reverse=True)
    anciennes = triees[:NOMBRE_BLOQUANTES]
    largeur = _largeur(presentation)
    _zone(
        slide,
        (MARGE, 1_250_000, largeur, 400_000),
        (
            f"{formater_entier(len(anciennes))} sur {formater_entier(resultat.total)}"
            " bloquantes ouvertes",
        ),
    )
    tableau = _tableau(
        slide,
        (MARGE, 1_750_000, largeur),
        EN_TETES_BLOQUANTES,
        [
            [
                anomalie.cle_ticket,
                anomalie.entite,
                libelle(anomalie.phase),
                anomalie.domaine,
                libelle(anomalie.statut),
                formater_entier(anomalie.age_jours_ouvres),
            ]
            for anomalie in anciennes
        ],
        (2, 2, 3, 3, 2, 2),
    )
    for ligne, anomalie in enumerate(anciennes, start=1):
        _lier(tableau.cell(ligne, 0), anomalie.cle_ticket, consultation)


def _annexe(
    presentation: Presentation_, consultation: Consultation, jour: date, contexte: ContexteRendu
) -> None:
    slide = _slide(presentation, DISPOSITION, TITRE_ANNEXE)
    largeur = _largeur(presentation)
    grille = contexte.grille
    codes = [] if grille is None else [c for c in CodeIndicateur if c in grille.indicateurs]
    _tableau(
        slide,
        (MARGE, 1_300_000, largeur),
        EN_TETES_GRILLE,
        [
            [
                code.value,
                libelle_indicateur(grille, code),
                formater_seuils(grille, code),
                contexte.version_grille,
            ]
            for code in codes
        ],
        (1, 6, 4, 2),
        dense=True,
    )
    _tableau(
        slide,
        (MARGE, 1_450_000 + HAUTEUR_LIGNE_DENSE * (len(codes) + 1), largeur),
        EN_TETES_TRACABILITE,
        [list(ligne) for ligne in lignes_tracabilite(consultation, jour, contexte)],
        (1, 3),
        dense=True,
    )


def rendre_powerpoint(
    consultation: Consultation, gabarit: bytes, date_arrete: date, contexte: ContexteRendu
) -> bytes:
    """Octets de la présentation du Passage à la Date d'arrêté, déterministes.

    Sans Météo à cette date : 6 slides, tableaux à en-têtes seuls, messages « aucune ».
    """
    meteos = consultation.meteo(date_arrete)
    presentation = _gabarit(gabarit)
    if len(presentation.slides) != 0:
        raise UsineErreur(
            f"gabarit PowerPoint : {len(presentation.slides)} slide(s), aucune attendue"
        )
    _titre(presentation, meteos, date_arrete, contexte)
    _synthese(presentation, consultation, meteos, date_arrete)
    _carte(presentation, consultation, meteos, date_arrete)
    for meteo in meteos:
        if meteo.couleur in COULEURS_SLIDE_ENTITE:
            _entite(presentation, consultation, meteo, date_arrete, contexte)
    _tendances(presentation, consultation)
    _bloquantes(presentation, consultation, date_arrete, contexte)
    _annexe(presentation, consultation, date_arrete, contexte)

    proprietes = presentation.core_properties
    proprietes.author = AUTEUR
    proprietes.last_modified_by = AUTEUR
    proprietes.title = titre_document(date_arrete)
    proprietes.revision = 1
    proprietes.created = horodatage_fixe(date_arrete)
    proprietes.modified = horodatage_fixe(date_arrete)
    proprietes.last_printed = horodatage_fixe(date_arrete)
    tampon = io.BytesIO()
    presentation.save(tampon)
    return archive_deterministe(tampon.getvalue(), date_arrete)
